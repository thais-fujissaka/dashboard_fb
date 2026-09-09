import streamlit as st
import pandas as pd
import numpy as np
import calendar
from datetime import datetime, timedelta
from utils.functions.general_functions_conciliacao import traduz_semana_mes, calcular_datas
from utils.functions.controladoria_planejamento_anual import insere_nova_linha
from utils.queries_cmv import *
from utils.queries_forecast import *
from utils.queries_dre_download import *
import openpyxl
# import os


pd.set_option('future.no_silent_downcasting', True)

datas = calcular_datas()

############################################ PROJEÇÕES MÊS CORRENTE ############################################

# Prepara df de faturamento agregado diário para a casa selecionada
def prepara_dados_faturam_agregado_diario(id_casa, df_faturamento_agregado_dia, fim_do_mes_atual, dois_meses_antes):
    # Filtra por casa
    df_faturamento_agregado_casa = df_faturamento_agregado_dia[
        (df_faturamento_agregado_dia['ID_Casa'] == id_casa) &
        df_faturamento_agregado_dia['Categoria'].isin(['Alimentos', 'Bebidas', 'Couvert', 'Serviço', 'Delivery', 'Gifts', 'Eventos A&B', 'Eventos Couvert', 'Eventos Locações', 'Outras Receitas'])
    ].copy()
    df_faturamento_agregado_casa['Data Evento'] = pd.to_datetime(df_faturamento_agregado_casa['Data Evento'], errors='coerce')
    
    # Traduz dia da semana para português
    df_faturamento_agregado_casa['Dia Semana'] = df_faturamento_agregado_casa['Data Evento'].dt.strftime('%A')
    df_faturamento_agregado_casa['Dia Semana'] = df_faturamento_agregado_casa['Dia Semana'].apply(
        lambda x: traduz_semana_mes(x, 'dia semana')
    )
    
    df_faturamento_agregado_casa['Dia_Mes'] = pd.to_datetime(df_faturamento_agregado_casa['Data Evento'], errors='coerce').dt.day

    # Filtra por casa e mês (anterior e corrente) - para utilizar no cálculo de projeção
    df_faturamento_agregado_mes_corrente = df_faturamento_agregado_casa[
        (df_faturamento_agregado_casa['Data Evento'] >= dois_meses_antes) &
        (df_faturamento_agregado_casa['Data Evento'] <= fim_do_mes_atual)
    ].copy()
    df_faturamento_agregado_mes_corrente = df_faturamento_agregado_mes_corrente.groupby(['ID_Casa', 'Casa', 'Categoria', 'Data Evento', 'Dia Semana', 'Dia_Mes'], as_index=False)[['Valor Bruto', 'Desconto', 'Valor Liquido']].sum()
    
    return df_faturamento_agregado_mes_corrente


# --- CRIA COMBINAÇÃO DE TODAS AS CATEGORIAS x DIAS (mês anterior e corrente) ---
def criar_df_dias(ano, mes):
    """Cria DataFrame com dias do mês, Data_Evento e Dia Semana traduzido."""
    ultimo_dia = calendar.monthrange(ano, mes)[1]
    dias = pd.DataFrame({'Dia_Mes': range(1, ultimo_dia + 1)})

    dias['Data Evento'] = pd.to_datetime({
        'year': ano,
        'month': mes,
        'day': dias['Dia_Mes']
    })

    dias['Dia Semana'] = dias['Data Evento'].dt.strftime('%A').apply(
        lambda x: traduz_semana_mes(x, 'dia semana')
    )
    return dias

def criar_df_dias_intervalo(ano_inicio, mes_inicio, ano_fim, mes_fim):
    datas = pd.period_range(
        start=f"{ano_inicio}-{mes_inicio:02d}",
        end=f"{ano_fim}-{mes_fim:02d}",
        freq="M"
    )

    lista_df = []

    for periodo in datas:
        ano = periodo.year
        mes = periodo.month
        lista_df.append(criar_df_dias(ano, mes))

    return pd.concat(lista_df, ignore_index=True)

def lista_dias_mes_anterior_atual(ano_atual, df_faturamento_agregado_mes_corrente):
    # Calcula faturamento mês corrente desde jan/2025
    ano_inicio = 2025
    mes_inicio = 1
    ano_fim = ano_atual + 1
    mes_fim = 12
    
    df_dias_mes = criar_df_dias_intervalo(
        ano_inicio=ano_inicio,
        mes_inicio=mes_inicio, 
        ano_fim=ano_fim,
        mes_fim=mes_fim
    )
    
    categorias = df_faturamento_agregado_mes_corrente['Categoria'].dropna().unique()
    df_categorias = pd.DataFrame({'Categoria': categorias})

    df_resultado = df_dias_mes.merge(df_categorias, how='cross')
    return df_resultado


# Função para cálculo da projeção - mês corrente
def cria_projecao_mes_corrente(df_faturamento_agregado_mes_corrente, df_dias_futuros_com_categorias):
    # Merge com df do mês completo para gerar projeção dos prox dias
    df_dias_futuros_mes = df_faturamento_agregado_mes_corrente.merge(
        df_dias_futuros_com_categorias, 
        how='right', 
        on=['Data Evento', 'Dia Semana', 'Categoria'])
    
    df_dias_futuros_mes['Faturamento Projetado'] = None

    # Loop por categoria
    for categoria in df_dias_futuros_mes['Categoria'].unique():
        df_cat = None
        if categoria not in ('Eventos A&B', 'Eventos Locações', 'Eventos Couvert', 'Outras Receitas'): 
            df_cat = df_dias_futuros_mes[df_dias_futuros_mes['Categoria'] == categoria].copy()
        
        if df_cat is not None and not df_cat.empty:
            for i, row in df_cat.iterrows():
                data = row['Data Evento']

                # -> estou fazendo para todos os dias (desde o mês anterior) para comparar projetado/real
                dia_semana = row['Dia Semana']

                # pega histórico das duas semanas anteriores MESMO DIA DA SEMANA e MESMA CATEGORIA
                duas_semanas_atras = data - timedelta(days=14)

                historico = df_dias_futuros_mes[
                    (df_dias_futuros_mes['Categoria'] == categoria) &
                    (df_dias_futuros_mes['Dia Semana'] == dia_semana) &
                    (df_dias_futuros_mes['Data Evento'] >= duas_semanas_atras) &
                    (df_dias_futuros_mes['Data Evento'] < data)
                ].copy()

                # usa o Valor_Bruto (real) quando existir, senão a Projeção (que pode vir de dias anteriores)
                # historico['Faturamento Projetado'] = historico['Faturamento Projetado'].fillna(0) # Para dias sem faturamento e sem projeção
                valores_para_media = historico['Valor Bruto'].fillna(historico['Faturamento Projetado']).astype(float)
                
                if not valores_para_media.empty:
                    media = valores_para_media.mean()
                    df_dias_futuros_mes.at[i, 'Faturamento Projetado'] = media
                    df_dias_futuros_mes['Faturamento Projetado'] = df_dias_futuros_mes['Faturamento Projetado'].fillna(0) # Para dias sem faturamento e sem projeção

    # Define valor final
    df_dias_futuros_mes["Valor Final"] = np.where(
        (df_dias_futuros_mes["Data Evento"].dt.month >= datas['mes_atual']) & (df_dias_futuros_mes['Data Evento'].dt.year == datas['ano_atual']),
        df_dias_futuros_mes["Faturamento Projetado"],       # usa o projetado se mês >= atual
        df_dias_futuros_mes["Valor Bruto"]            # senão usa o real
    )
    # df_dias_futuros_mes['Valor Final'] = df_dias_futuros_mes['Valor Bruto'].fillna(df_dias_futuros_mes['Faturamento Projetado'])
    return df_dias_futuros_mes


def aplica_layout_mes_corrente(df_dias_futuros_mes, df_faturamento_eventos, df_parc_receit_extr, df_dias_mes, id_casa, casa, mes_selecionado, ano_selecionado):
    # Prepara dados de faturamentos
    df_dias_futuros_mes_filtrado = df_dias_futuros_mes.copy()
    df_dias_futuros_mes_filtrado['ID_Casa'] = df_dias_futuros_mes_filtrado['ID_Casa'].fillna(id_casa)
    df_dias_futuros_mes_filtrado['Casa'] = df_dias_futuros_mes_filtrado['Casa'].fillna(casa)

    df_dias_futuros_mes_filtrado = df_dias_futuros_mes_filtrado[df_dias_futuros_mes_filtrado['Categoria'].isin(['Alimentos', 'Bebidas', 'Couvert', 'Delivery', 'Gifts'])].copy()
    df_dias_futuros_mes_filtrado = df_dias_futuros_mes_filtrado[['Categoria', 'Data Evento', 'Dia Semana', 'Valor Final']]
    df_dias_futuros_mes_filtrado = df_dias_futuros_mes_filtrado.rename(columns={'Valor Final': 'Valor Projetado'})
    
    # Prepada dados de Eventos
    df_faturamento_eventos_filtrado = df_faturamento_eventos[df_faturamento_eventos['ID_Casa'] == id_casa].copy()
    df_faturamento_eventos_filtrado = df_faturamento_eventos_filtrado[['Categoria', 'Data Evento', 'Valor Bruto']]
    df_faturamento_eventos_filtrado = df_faturamento_eventos_filtrado.rename(columns={'Valor Bruto': 'Valor Projetado'})

    # Prepara dados de Receitas Extraordinárias
    df_faturamento_rec_extr_filtrado = df_parc_receit_extr[(df_parc_receit_extr['ID_Casa'] == id_casa) & (df_parc_receit_extr['Categoria'] == 'Outras Receitas')].copy()
    df_faturamento_rec_extr_filtrado = df_faturamento_rec_extr_filtrado[['Categoria', 'Data Evento', 'Valor Bruto']]
    df_faturamento_rec_extr_filtrado = df_faturamento_rec_extr_filtrado.rename(columns={'Valor Bruto': 'Valor Projetado'})
    
    # Concatena os dados
    df_concat = pd.concat([df_dias_futuros_mes_filtrado, df_faturamento_eventos_filtrado, df_faturamento_rec_extr_filtrado])
    
    df_todos_dias_mes = pd.merge(
        df_dias_mes,
        df_concat,
        on=['Data Evento', 'Categoria'],
        how='left'
    )
    
    # Cria coluna de dia da semana para cada dia
    df_todos_dias_mes['Dia Semana'] = df_todos_dias_mes['Data Evento'].dt.strftime('%A')
    df_todos_dias_mes['Dia Semana'] = df_todos_dias_mes['Dia Semana'].apply(
        lambda x: traduz_semana_mes(x, 'dia semana')
    )
    df_todos_dias_mes['Valor Projetado'] = df_todos_dias_mes['Valor Projetado'].fillna(0)
    
    # Filtra para mês/ano corrente
    df_todos_dias_mes_corrente = df_todos_dias_mes[(df_todos_dias_mes['Data Evento'].dt.month == mes_selecionado) & (df_todos_dias_mes['Data Evento'].dt.year == ano_selecionado)].copy()

    df_todos_dias_mes_corrente['Data Evento'] = df_todos_dias_mes_corrente['Data Evento'].dt.strftime('%d/%m/%Y') 
    
    # Evitar duplicados Data X Categoria antes de pivotar
    df_todos_dias_mes_corrente = df_todos_dias_mes_corrente.groupby(['Data Evento', 'Dia Semana', 'Categoria'], as_index=False)['Valor Projetado'].sum()
    
    # Categorias viram colunas
    pivot_faturamento_mes_corrente = df_todos_dias_mes_corrente.pivot(
        index=['Data Evento', 'Dia Semana'],
        columns='Categoria',
        values='Valor Projetado'
    ).reset_index()
    
    # Coluna para Total de cada dia
    colunas_para_total = [
        'Alimentos',
        'Bebidas',
        'Couvert',
        'Delivery',
        'Gifts',
        'Eventos A&B',
        'Eventos Locações',
        'Eventos Couvert',
        'Outras Receitas'
    ]

    # pega só as colunas que existem no dataframe
    colunas_existentes = [
        col for col in colunas_para_total if col in pivot_faturamento_mes_corrente.columns
    ]

    pivot_faturamento_mes_corrente['Total'] = (
        pivot_faturamento_mes_corrente[colunas_existentes]
        .astype(float)
        .sum(axis=1)
    )

    return pivot_faturamento_mes_corrente


# Destaca dias futuros do mês corrente
def destaca_dias_futuros_mes_corrente(row):
    hoje = datas['today'].strftime('%d/%m/%Y')
    estilos = [''] * len(row)

    # regra linha inteira (dias futuros)
    if row['Data Evento'] >= hoje:
        estilos = ['background-color: rgba(255,255,224);'] * len(row)

    # regra coluna TOTAL 
    if 'Total' in row.index:
        idx_total = row.index.get_loc('Total')
        if row['Data Evento'] >= hoje:
            estilos[idx_total] = 'background-color: rgba(255,255,224); color: black; font-weight: 500;'
        else:
            estilos[idx_total] = 'color: black; font-weight: 500;'

    return estilos


############################################ PROJEÇÕES - PRÓXIMOS MESES ############################################

# Une faturamentos e orçamentos mensais para calcular histórico de atingimento (%)
def prepara_dados_faturamento_orcamentos_mensais(id_casa, df_orcamentos, df_faturamento_agregado_mes, df_demais_receitas_extr, df_bilheterias, df_ajustes_manuais, ano_passado, ano_atual, ano_selecionado):
    # Filtra por casa e período (ano passado e atual)
    df_orcamentos_casa = df_orcamentos[
        (df_orcamentos['ID_Casa'] == id_casa) &
        (df_orcamentos['Ano'] >= ano_passado) &
        (df_orcamentos['Ano'] <= ano_atual) &
        (df_orcamentos['Classificacao_Contabil_1'] == 'Faturamento Bruto')
    ].copy()
    
    df_faturamento_mes_casa = df_faturamento_agregado_mes[
        (df_faturamento_agregado_mes['ID_Casa'] == id_casa) &
        (df_faturamento_agregado_mes['Ano'] >= ano_passado) &
        (df_faturamento_agregado_mes['Ano'] <= ano_atual)
    ].copy()

    if id_casa in [110, 128, 145]: # Apenas casas que tem Bilheteria
        if id_casa == 110: # Blue Note
            # Caso - Abril
            df_faturamento_mes_casa = df_faturamento_mes_casa.drop_duplicates(subset=['ID_Casa', 'Casa', 'Categoria', 'Ano', 'Mês'], keep='first')
            # Bilheteria vem das Receitas Extraordinárias
            df_bilheteria = df_demais_receitas_extr[(df_demais_receitas_extr['Casa'] == 'Blue Note - São Paulo') & (df_demais_receitas_extr['Classificacao'] == 'Bilheteria')].copy()
            df_bilheteria['Mês'] = df_bilheteria['Data_Ocorrencia'].dt.month
            df_bilheteria['Ano'] = df_bilheteria['Data_Ocorrencia'].dt.year
            df_bilheteria = df_bilheteria.groupby(['Casa', 'Ano', 'Mês'], as_index=False)['Valor Bruto'].sum()
            df_bilheteria['Categoria'] = 'Couvert'
            df_bilheteria = df_bilheteria.rename(columns={'Valor Bruto': 'Valor Bilheteria'})
            df_faturamento_mes_casa = pd.merge(df_faturamento_mes_casa, df_bilheteria, on=['Casa', 'Categoria', 'Mês', 'Ano'], how='left').fillna(0)
            df_faturamento_mes_casa['Valor Bilheteria'] = pd.to_numeric(df_faturamento_mes_casa['Valor Bilheteria'], errors='coerce')
            df_faturamento_mes_casa['Valor Bruto'] += df_faturamento_mes_casa['Valor Bilheteria']
            
        else: # Love Cabaret, Ultra Evil
            df_bilheteria = df_bilheterias[df_bilheterias['ID_Casa'] == id_casa].copy()
            df_bilheteria['Categoria'] = 'Couvert'
            df_bilheteria = df_bilheteria.groupby(['ID_Casa', 'Casa', 'Categoria', 'Ano', 'Mês'], as_index=False)[['Valor Bruto', 'Desconto', 'Valor Liquido']].sum()
            df_faturamento_mes_casa = pd.concat([df_faturamento_mes_casa, df_bilheteria])
    
    df_faturamento_mes_casa['Valor Bruto'] = pd.to_numeric(df_faturamento_mes_casa['Valor Bruto'], errors='coerce')
    df_faturamento_mes_casa = df_faturamento_mes_casa.groupby(['ID_Casa', 'Casa', 'Categoria', 'Ano', 'Mês'], as_index=False)['Valor Bruto'].sum()

    # Inclui ajustes manuais para itens de faturamento que tem lançamento de ajuste
    df_ajustes_categoria = df_ajustes_manuais[
        (df_ajustes_manuais['ID_Casa'] == id_casa) &
        (df_ajustes_manuais['Ano'] == ano_selecionado) &
        (df_ajustes_manuais['Classificacao_Contabil_1'] == 'Faturamento Bruto')
    ].copy()

    if not df_ajustes_categoria.empty:
        df_ajustes_categoria = df_ajustes_categoria.groupby(['ID_Casa', 'Casa', 'Mês', 'Ano', 'Classificacao_Contabil_1', 'Classificacao_Contabil_2'], as_index=False)['Valor Ajuste'].sum()
        df_ajustes_categoria = df_ajustes_categoria[['ID_Casa', 'Casa', 'Classificacao_Contabil_2', 'Mês', 'Ano', 'Valor Ajuste']]
        df_ajustes_categoria = df_ajustes_categoria.rename(columns={'Classificacao_Contabil_2': 'Categoria', 'Valor Ajuste': 'Valor Bruto'})
        df_ajustes_categoria['Categoria'] = df_ajustes_categoria['Categoria'].replace({
            'Alimentação': 'Alimentos',
            'Bebida': 'Bebidas',
            'Artístico (couvert/shows)': 'Couvert',
        })
        df_faturamento_mes_casa = pd.concat([df_faturamento_mes_casa, df_ajustes_categoria])
        df_faturamento_mes_casa['Valor Bruto'] = pd.to_numeric(df_faturamento_mes_casa['Valor Bruto'], errors='coerce')
        df_faturamento_mes_casa = df_faturamento_mes_casa.groupby(['ID_Casa', 'Casa', 'Categoria', 'Mês', 'Ano'], as_index=False)['Valor Bruto'].sum()
        
    if id_casa in [149, 110]:
        if id_casa == 149: # Priceless - Eventos Rebate Fornecedores
            df_receitas_extr = df_demais_receitas_extr[(df_demais_receitas_extr['Casa'] == 'Priceless') & (df_demais_receitas_extr['Cliente'] == 'TORANJA ')].copy()
            cols_agrupa = ['Casa', 'Cliente', 'Classificacao', 'Mês', 'Ano']
            
        if id_casa == 110: # Blue Note - Itens Fechamento Financeiro que vão para Faturamento
            df_receitas_extr = df_demais_receitas_extr[(df_demais_receitas_extr['Casa'] == 'Blue Note - São Paulo') & (df_demais_receitas_extr['Classificacao'].isin(['Membership', 'Premium Corp']))].copy()
            cols_agrupa = ['Casa', 'Classificacao', 'Mês', 'Ano']
        
        df_receitas_extr['Mês'] = df_receitas_extr['Data_Ocorrencia'].dt.month
        df_receitas_extr['Ano'] = df_receitas_extr['Data_Ocorrencia'].dt.year
        df_receitas_extr = df_receitas_extr.drop(columns=['Valor Bruto'])
        df_receitas_extr = df_receitas_extr.rename(columns={'Valor_Total': 'Valor Bruto'}) # A VERIFICAR # 

        df_receitas_extr = df_receitas_extr.groupby(cols_agrupa, as_index=False)['Valor Bruto'].sum()
        df_receitas_extr['ID_Casa'] = id_casa
        df_receitas_extr = df_receitas_extr.rename(columns={'Classificacao': 'Categoria'})
        df_receitas_extr = df_receitas_extr[['ID_Casa', 'Casa', 'Categoria', 'Mês', 'Ano', 'Valor Bruto']]
        df_faturamento_mes_casa = pd.concat([df_faturamento_mes_casa, df_receitas_extr])

        df_faturamento_mes_casa['Valor Bruto'] = pd.to_numeric(df_faturamento_mes_casa['Valor Bruto'], errors='coerce')
        df_faturamento_mes_casa = df_faturamento_mes_casa.groupby(['ID_Casa', 'Casa', 'Categoria', 'Mês', 'Ano'], as_index=False)['Valor Bruto'].sum()
        df_faturamento_mes_casa['Categoria'] = df_faturamento_mes_casa['Categoria'].replace({
            'Eventos': 'Eventos Rebate Fornecedores',
            'Premium Corp': 'Eventos Rebate Fornecedores'
        })

    # Merge para calcular faturamento/orçamento
    df_faturamento_orcamento = pd.merge(
        df_faturamento_mes_casa[['Categoria', 'Ano', 'Mês', 'Valor Bruto']],
        df_orcamentos_casa[['Categoria', 'Ano', 'Mês', 'Orçamento']],
        how='right',
        on=['Mês', 'Ano', 'Categoria']
    )
    
    # Calcula Faturamento / Orçamento
    df_faturamento_orcamento['Atingimento Real'] = (
        (df_faturamento_orcamento['Valor Bruto'].astype(float) /
        df_faturamento_orcamento['Orçamento'].replace(0, np.nan).astype(float)) * 100
    )

    return df_faturamento_mes_casa, df_faturamento_orcamento


# Cria combinação das categorias de faturamento com meses do ano (desde 2025)
def lista_meses_ano(lista_itens_faturamento):
    ano_inicial = 2025
    ano_atual = datetime.now().year

    # Lista de anos (2025 até ano atual)
    anos = list(range(ano_inicial, ano_atual + 1))
    meses = list(range(1, 13))

    # Cria todas combinações de Ano x Mês
    df_meses_anos = pd.DataFrame(
        [(ano, mes) for ano in anos for mes in meses],
        columns=['Ano', 'Meses_Ano']
    )

    # Cross com categorias
    df_categorias = pd.DataFrame({'Categoria': lista_itens_faturamento})

    df_final = df_meses_anos.merge(df_categorias, how='cross')

    # Cria coluna de data (primeiro dia do mês)
    df_final['Data'] = pd.to_datetime({
        'year': df_final['Ano'],
        'month': df_final['Meses_Ano'],
        'day': 1
    })

    return df_final


# Função para cálculo da projeção - meses seguintes
def projecao_faturamento_meses_seguintes(df_faturamento_orcamento, df_meses_futuros_com_categorias, ano_atual, mes_atual):
    # Merge com df que contém todos os meses (ano anterior e corrente)
    df_meses_seguintes = df_faturamento_orcamento.merge(
        df_meses_futuros_com_categorias, 
        how='right', 
        left_on=['Ano', 'Mês', 'Categoria'],
        right_on=['Ano', 'Meses_Ano', 'Categoria']
    )
    df_meses_seguintes = df_meses_seguintes[df_meses_seguintes['Categoria'].isin(['Alimentos', 'Bebidas', 'Couvert', 'Delivery', 'Eventos A&B', 'Eventos Couvert', 'Eventos Locações', 'Eventos Rebate Fornecedores', 'Gifts', 'Membership', 'Outras Receitas', 'Serviço'])].copy()
    df_meses_seguintes['Projeção Atingimento'] = None
    df_meses_seguintes['Valor Projetado'] = None
    
    # Loop por categoria
    for categoria in df_meses_seguintes['Categoria'].unique():
        if categoria != 'Serviço':
            df_cat = None
            df_cat = df_meses_seguintes[df_meses_seguintes['Categoria'] == categoria].copy()
            
            if df_cat is not None and not df_cat.empty:
                for i, row in df_cat.iterrows():
                    mes = row['Data']

                    # pega histórico dos dois meses atrás
                    dois_meses_atras = mes - pd.DateOffset(months=2)

                    historico = df_meses_seguintes[
                        (df_meses_seguintes['Categoria'] == categoria) &
                        (df_meses_seguintes['Data'] >= dois_meses_atras) &
                        (df_meses_seguintes['Data'] < mes)
                    ].copy()

                    # Define colunas auxiliares conforme o mês
                    historico["Atingimento_Usado"] = np.where(
                        (historico["Mês"] >= mes_atual) & (historico['Ano'] == ano_atual),
                        historico["Projeção Atingimento"],       # usa o projetado se mês >= atual
                        historico["Atingimento Real"]            # senão usa o real
                    )
                    
                    # usa o Atingimento (real) quando existir, senão a Projecao (que pode vir de meses anteriores)
                    valores_para_media = historico['Atingimento_Usado'].fillna(historico['Projeção Atingimento']).astype(float)

                    if not valores_para_media.empty:
                        media = valores_para_media.mean()
                        df_meses_seguintes.at[i, 'Projeção Atingimento'] = media

    # Define valor de faturamento projetado baseado na projeção (%) de atingimento do orçamento
    df_meses_seguintes['Valor Projetado'] = (
        df_meses_seguintes['Orçamento'].astype(float) * (df_meses_seguintes['Projeção Atingimento'].astype(float) / 100)
    )
    return df_meses_seguintes


# Função para cálculo da projeção do serviço - meses seguintes: 13% do faturamento A&B projetado
# Precisa dos faturamentos das outras categorias já calculado
def projecao_faturamento_servico_meses_seguintes(df_faturamento_meses_futuros, ano_atual, mes_atual):
    # Filtra o df de faturamento para apenas a categoria de Serviço
    df_cat = df_faturamento_meses_futuros[df_faturamento_meses_futuros['Categoria'] == 'Serviço']
    if df_cat is not None and not df_cat.empty:
        for i, row in df_cat.iterrows():
            mes = row['Mês']
            ano = row['Ano']

            historico = df_faturamento_meses_futuros[
                (df_faturamento_meses_futuros['Categoria'].isin(['Alimentos', 'Bebidas'])) &
                (df_faturamento_meses_futuros['Mês'] == mes) &
                (df_faturamento_meses_futuros['Ano'] == ano)
            ].copy()

            # Define colunas auxiliares conforme o mês
            historico["Faturamento_Usado"] = np.where(
                (historico["Mês"] >= mes_atual) & (historico['Ano'] == ano_atual),
                historico["Valor Projetado"],       # usa o projetado se mês >= atual
                historico["Valor Bruto"]            # senão usa o real
            )
            
            # usa o Atingimento (real) quando existir, senão a Projecao (que pode vir de meses anteriores)
            historico['Faturamento_Usado'] = pd.to_numeric(historico['Faturamento_Usado'], errors='coerce')
            soma_ab = historico['Faturamento_Usado'].sum()
            valor_servico = soma_ab * 0.13
            
            df_faturamento_meses_futuros.at[i, 'Projeção Atingimento'] = '-'
            df_faturamento_meses_futuros.at[i, 'Valor Projetado'] = valor_servico

    return df_faturamento_meses_futuros


def projecao_impostos_venda(df_faturamento_para_impostos, lista_itens_impostos, df_impostos_meses_futuros, df_parametros_impostos, casa):
    df_final = df_impostos_meses_futuros.copy() # Df com lista com meses futuros
    df_final = df_final.rename(columns={'Meses_Ano': 'Mês'})

    for item in lista_itens_impostos:
        if item == 'ISS':
            categorias_fat_para_soma = ['Eventos Couvert', 'Eventos Locações', 'Eventos Rebate Fornecedores', 'Gifts']
        elif item == 'ICMS':
            categorias_fat_para_soma = ['Eventos A&B', 'Alimentos', 'Bebidas', 'Couvert', 'Delivery']
        elif item == 'PIS':
            categorias_fat_para_soma = ['Alimentos', 'Bebidas', 'Couvert', 'Delivery', 'Eventos A&B', 'Eventos Couvert', 'Eventos Locações', 'Eventos Rebate Fornecedores', 'Gifts']
        elif item == 'COFINS':
            categorias_fat_para_soma = ['Alimentos', 'Bebidas', 'Couvert', 'Delivery', 'Eventos A&B', 'Eventos Couvert', 'Eventos Locações', 'Eventos Rebate Fornecedores', 'Gifts']

        # Calcula o valor do imposto a partir dos itens de faturamento e porcentagem correspondente
        df_imposto = df_faturamento_para_impostos[df_faturamento_para_impostos['Categoria'].isin(categorias_fat_para_soma)].copy()
        df_imposto["Valor"] = np.where(
            (df_imposto["Mês"] >= datas['mes_atual']) & (df_imposto['Ano'] == datas['ano_atual']),
            df_imposto["Valor Projetado"],       # usa o projetado se mês >= atual
            df_imposto["Valor Bruto"]            # senão usa o real
        )
        df_imposto['Valor'] = pd.to_numeric(df_imposto['Valor'], errors='coerce')
        df_imposto = df_imposto.groupby(['Ano', 'Mês', 'Data'], as_index=False)['Valor'].sum()
        df_imposto['Imposto'] = item

        # Aplica taxas de acordo com as faixas de data de vigência
        df_imposto = df_imposto.merge(df_parametros_impostos[['Data Inicio', 'Data Fim', 'Taxa', 'Classificacao_Contabil_2']], how='cross')
        df_imposto = df_imposto[
            ((df_imposto['Data Fim'].notna()) &
            (df_imposto['Data'] >= df_imposto['Data Inicio']) &
            (df_imposto['Data'] <= df_imposto['Data Fim']) &
            (df_imposto['Imposto'] == df_imposto['Classificacao_Contabil_2']))
            |
            ((df_imposto['Data Fim'].isna()) & # Caso a data de fim de vigência seja nula (vigência atual)
            (df_imposto['Data'] >= df_imposto['Data Inicio']) &
            (df_imposto['Imposto'] == df_imposto['Classificacao_Contabil_2']))]

        if casa == 'Bar Léo - Centro' and item == 'ICMS': # Caso específico
            df_imposto[f'Valor {item}'] = 0
        else:
            df_imposto['Taxa'] = pd.to_numeric(df_imposto['Taxa'])
            df_imposto[f'Valor {item}'] = df_imposto['Valor'] * (df_imposto['Taxa'] / 100) 

        # Dependem do ICMS
        if item in ['PIS', 'COFINS']:
            df_icms = df_final.drop_duplicates(subset=['Ano', 'Mês', 'Data'])
            df_imposto = pd.merge(
                df_imposto,
                df_icms[['Ano', 'Mês', 'Data', 'Valor ICMS']],
                on=['Ano', 'Mês', 'Data'],
                how='left'
            )
            df_imposto[f'Valor {item}'] = (df_imposto['Valor'] - df_imposto['Valor ICMS']) * (df_imposto['Taxa'] / 100) 

        df_final = df_final.merge(
            df_imposto[['Ano', 'Mês', 'Data', f'Valor {item}']],
            on=['Ano', 'Mês', 'Data'],
            how='left'
        )
    df_final['Valor PIS / COFINS'] = df_final['Valor PIS'] + df_final['Valor COFINS']
    
    # Transforma: cada coluna era o valor de um imposto -> apenas uma de 'Valor'
    df_long = df_final.melt(
        id_vars=['Ano', 'Mês', 'Data'],
        value_vars=['Valor ISS', 'Valor ICMS', 'Valor PIS', 'Valor COFINS', 'Valor PIS / COFINS'],
        var_name='Categoria',
        value_name='Valor'
    )
    df_long['Categoria'] = df_long['Categoria'].str.replace('Valor ', '')
    df_long = df_long.drop_duplicates()
    return df_long


def projecao_impostos_renda(df_faturamento_para_impostos, lista_itens_impostos, df_impostos_meses_futuros, df_parametros_impostos, casa):
    df_parametros_impostos_renda = df_parametros_impostos[df_parametros_impostos['Classificacao_Contabil_1'] == 'Imposto de Renda'].copy()
    df_final = df_impostos_meses_futuros.copy() # Df com lista com meses futuros
    df_final = df_final.rename(columns={'Meses_Ano': 'Mês'})

    for item in lista_itens_impostos:
        if item in ['CSLL s/ Serviço', 'IRPJ s/ Serviço']:
            categorias_fat_para_soma = ['Eventos Couvert', 'Eventos Locações', 'Eventos Rebate Fornecedores', 'Gifts']
        elif item in ['CSLL s/ Comercio', 'IRPJ s/ Comercio']:
            categorias_fat_para_soma = ['Eventos A&B', 'Alimentos', 'Bebidas', 'Couvert', 'Delivery']
        elif item in ['CSLL Não operacional', 'IRPJ Não operacional']:
            categorias_fat_para_soma = ['Outras Receitas']
        
        # Calcula o valor do imposto a partir dos itens de faturamento e porcentagem correspondente
        df_imposto = df_faturamento_para_impostos[df_faturamento_para_impostos['Categoria'].isin(categorias_fat_para_soma)].copy()
        df_imposto["Valor"] = np.where(
            (df_imposto["Mês"] >= datas['mes_atual']) & (df_imposto['Ano'] == datas['ano_atual']),
            df_imposto["Valor Projetado"],       # usa o projetado se mês >= atual
            df_imposto["Valor Bruto"]            # senão usa o real
        )
        df_imposto['Valor'] = pd.to_numeric(df_imposto['Valor'], errors='coerce')
        df_imposto = df_imposto.groupby(['Ano', 'Mês', 'Data'], as_index=False)['Valor'].sum()
        df_imposto['Imposto'] = item

        # Aplica taxas de acordo com as faixas de data de vigência
        df_imposto = df_imposto.merge(df_parametros_impostos_renda[['Data Inicio', 'Data Fim', 'Taxa', 'Classificacao_Contabil_2']], how='cross')
        df_imposto = df_imposto[
            ((df_imposto['Data Fim'].notna()) &
            (df_imposto['Data'] >= df_imposto['Data Inicio']) &
            (df_imposto['Data'] <= df_imposto['Data Fim'])) &
            (df_imposto['Imposto'] == df_imposto['Classificacao_Contabil_2'])
            |
            ((df_imposto['Data Fim'].isna()) & # Caso a data de fim de vigência seja nula (vigência atual)
            (df_imposto['Data'] >= df_imposto['Data Inicio']) &
            (df_imposto['Imposto'] == df_imposto['Classificacao_Contabil_2']))]
        
        df_imposto['Taxa'] = pd.to_numeric(df_imposto['Taxa'])
        df_imposto[f'Valor {item}'] = df_imposto['Valor'] * (df_imposto['Taxa'] / 100) 
        if item in ['CSLL s/ Serviço', 'CSLL s/ Comercio', 'CSLL Não operacional']: # Caso específico - itens CSLL
            df_imposto[f'Valor {item}'] *= 0.09 
        
        df_final = df_final.merge(
            df_imposto[['Ano', 'Mês', 'Data', f'Valor {item}']],
            on=['Ano', 'Mês', 'Data'],
            how='left'
        )
    # Calcula valores totais de IRPJ e CSLL
    df_final['Valor IRPJ Total'] = (df_final['Valor IRPJ s/ Serviço'] + df_final['Valor IRPJ s/ Comercio'] + df_final['Valor IRPJ Não operacional']) * 0.15
    df_final['Valor CSLL Total'] = (df_final['Valor CSLL s/ Serviço'] + df_final['Valor CSLL s/ Comercio'] + df_final['Valor CSLL Não operacional'])

    # Base para cálculo do adicional de IRPJ 10%
    df_final['Base IRPJ'] = (df_final['Valor IRPJ s/ Serviço'] + df_final['Valor IRPJ s/ Comercio'] + df_final['Valor IRPJ Não operacional'])
    
    # Identifica trimestre
    df_final['Trimestre'] = ((df_final['Data'].dt.month - 1) // 3) + 1
    df_final['Trimestre'] = df_final['Trimestre'].astype('Int64')
    # Soma da base por ano/trimestre
    base_trimestre = df_final[df_final['Categoria'].isin(['IRPJ s/ Serviço', 'IRPJ s/ Comercio', 'IRPJ Não operacional'])]
    base_trimestre = base_trimestre.drop_duplicates(subset=['Ano', 'Mês'], keep='first')
    base_trimestre = base_trimestre.groupby(['Ano', 'Trimestre'])['Base IRPJ'].sum().reset_index(name='Base Trimestre')
    
    # Calcula 10%
    base_trimestre['IRPJ 10%'] = np.where(
        base_trimestre['Base Trimestre'] >= 60000,
        (base_trimestre['Base Trimestre'] - 60000) * 0.10, # HARDCODED #
        0
    )
    # Junta no dataframe principal
    df_final = df_final.merge(
        base_trimestre[['Ano', 'Trimestre', 'IRPJ 10%']],
        on=['Ano', 'Trimestre'],
        how='left'
    )
    df_final['Valor IRPJ 10%'] = 0.0 # Zera todos os meses

    # Aplica apenas no último mês de cada trimestre
    df_final.loc[df_final['Mês'].isin([3, 6, 9, 12]), 'Valor IRPJ 10%'] = df_final.loc[df_final['Mês'].isin([3, 6, 9, 12]), 'IRPJ 10%']

    # Transforma: cada coluna era o valor de um imposto -> apenas uma de 'Valor'
    df_long = df_final.melt(
        id_vars=['Ano', 'Mês', 'Data'],
        value_vars=['Valor IRPJ Total', 'Valor CSLL Total', 'Valor IRPJ 10%'], # Desconsidera as outras colunas de valores que compõem os totais
        var_name='Categoria',
        value_name='Valor'
    )
    df_long['Categoria'] = df_long['Categoria'].replace({
        'Valor IRPJ Total': 'IRPJ', 
        'Valor CSLL Total': 'CSLL',
        'Valor IRPJ 10%': 'IRPJ 10%',
    })
    df_long = df_long.drop_duplicates()
    return df_long


def projecao_imposto_simples(df_gorjeta, df_salarios, df_faturamento, df_aliquotas_imp_simples, df_impostos_renda_dre, mes_selecionado, ano_selecionado, casa):
    data_atual = pd.Timestamp(year=datas['ano_atual'], month=datas['mes_atual'], day=1)
    
    if casa == 'Bar Léo - Centro': # Calcula imposto com base no Faturamento Bruto
        df_faturamento = df_faturamento.groupby(['Ano', 'Mês', 'Data'], as_index=False)[['Valor Bruto', 'Valor Projetado']].sum()
        df_faturamento['Valor'] = np.where(
            df_faturamento['Data'] >= data_atual,
            df_faturamento['Valor Projetado'],
            df_faturamento['Valor Bruto']
        )
        df_base_imposto = df_faturamento[['Data', 'Valor']]
    
    elif casa == 'Blue Note - São Paulo': # Calcula imposto sobre (Alimentos + Bebidas) * 25%
        df_faturamento = df_faturamento[df_faturamento['Categoria'].isin(['Alimentos', 'Bebidas'])].copy()
        df_faturamento = df_faturamento.groupby(['Ano', 'Mês', 'Data'], as_index=False)[['Valor Bruto', 'Valor Projetado']].sum()
        df_faturamento['Valor'] = np.where(
            df_faturamento['Data'] >= data_atual,
            df_faturamento['Valor Projetado'],
            df_faturamento['Valor Bruto']
        )
        df_faturamento['Valor'] *= 0.25
        df_base_imposto = df_faturamento[['Data', 'Valor']]

    elif casa in ['Bar Brahma - Centro', 'Bar Brahma - Granja', 'Bar Brahma - Paulista', 'Girondino', 'Jacaré', 'Orfeu', 'Priceless', 'Riviera Bar']: # Calcula imposto com base em Gorjeta + Salários
        df_gorjeta['Valor'] = np.where(
            df_gorjeta['Data'] >= data_atual,
            df_gorjeta['Custo Projetado'],
            df_gorjeta['Custo Real']
        )
        df_salarios['Valor'] = np.where(
            df_salarios['Data'] >= data_atual,
            df_salarios['Custo Projetado'],
            df_salarios['Custo Real']
        )
        df_base_imposto = pd.concat([df_gorjeta[['Data', 'Valor']], df_salarios[['Data', 'Valor']]])
        df_base_imposto = df_base_imposto.groupby('Data', as_index=False)['Valor'].sum() # Soma gorjeta e salários por mês
    
    primeiro_ano = df_base_imposto['Data'].dt.year.min() # Primeiro ano da projeção (2025)
    valor_base = df_base_imposto.loc[ # Pega o acumulado de janeiro do ano seguinte
        (df_base_imposto['Data'].dt.year == primeiro_ano + 1)
        & (df_base_imposto['Data'].dt.month == 1),
        'Valor'
    ].iloc[0]

    df_base_imposto.loc[df_base_imposto['Data'].dt.year == primeiro_ano, 'Valor'] = valor_base # Aplica em todos os meses do primeiro ano
    df_base_imposto['Acumulado_12M'] = df_base_imposto['Valor'].shift(1).rolling(12, min_periods=12).sum() # Acumulado de cada mês - soma os 12 meses anteriores
    
    # Aplica data de vigência e alíquota de acordo com acumulado de cada mês
    df_imposto = df_base_imposto.merge(df_aliquotas_imp_simples, how='cross')
    df_imposto['MINIMO_RECEITA_BRUTA'] = pd.to_numeric(df_imposto['MINIMO_RECEITA_BRUTA'], errors='coerce')
    df_imposto['MAXIMO_RECEITA_BRUTA'] = pd.to_numeric(df_imposto['MAXIMO_RECEITA_BRUTA'], errors='coerce')
    df_imposto['ALIQUOTA'] = pd.to_numeric(df_imposto['ALIQUOTA'], errors='coerce')
    df_imposto['VALOR_DEDUZIR'] = pd.to_numeric(df_imposto['VALOR_DEDUZIR'], errors='coerce')
    df_imposto['Valor'] = pd.to_numeric(df_imposto['Valor'], errors='coerce')
    # st.write(df_imposto)
    df_imposto = df_imposto[
        ((df_imposto['Data Fim'].notna()) &
        (df_imposto['Data'] >= df_imposto['Data Inicio']) &
        (df_imposto['Data'] <= df_imposto['Data Fim']) &
        (df_imposto['Acumulado_12M'] >= df_imposto['MINIMO_RECEITA_BRUTA']) &
        (df_imposto['Acumulado_12M'] <= df_imposto['MAXIMO_RECEITA_BRUTA'])) 
        |
        ((df_imposto['Data Fim'].isna()) & # Caso a data de fim de vigência seja nula (vigência atual)
        (df_imposto['Data'] >= df_imposto['Data Inicio']) &
        (df_imposto['Acumulado_12M'] >= df_imposto['MINIMO_RECEITA_BRUTA']) &
        (df_imposto['Acumulado_12M'] <= df_imposto['MAXIMO_RECEITA_BRUTA']))]
    
    # Calcula alíquota real e valor do imposto por mês
    df_imposto['Aliquota Real'] = ((df_imposto['Acumulado_12M'] * (df_imposto['ALIQUOTA'] / 100)) - df_imposto['VALOR_DEDUZIR']) / df_imposto['Acumulado_12M']
    df_imposto['Valor Imposto'] = df_imposto['Aliquota Real'] * df_imposto['Valor']

    if casa == 'Bar Léo - Centro': # Calcula valor recolhimento ICMS
        df_imposto['Valor Recolhimento'] = df_imposto['Valor'] * 0.04 # HARDCODED # TAXA DO ICMS #
        df_imposto['Valor Imposto'] = df_imposto['Valor Imposto'] + df_imposto['Valor Recolhimento']
    # st.write('simples', df_imposto)
    # Incui valor do SIMPLES do layout de impostos de renda
    df_imposto_simples_dre = df_imposto[['Data', 'Valor Imposto']]
    df_imposto_simples_dre['Categoria'] = 'SIMPLES'
    df_imposto_simples_dre = df_imposto_simples_dre[(df_imposto_simples_dre['Data'].dt.month == mes_selecionado) & (df_imposto_simples_dre['Data'].dt.year == ano_selecionado)]

    df_impostos_renda_dre = pd.merge(
        df_impostos_renda_dre,
        df_imposto_simples_dre[['Categoria', 'Valor Imposto']],
        on=['Categoria'],
        how='left'
    )
    df_impostos_renda_dre.loc[df_impostos_renda_dre['Categoria'] == 'SIMPLES', 'Valor Real'] += df_impostos_renda_dre['Valor Imposto']
    df_impostos_renda_dre = df_impostos_renda_dre.drop(columns=['Valor Imposto'])
    return df_impostos_renda_dre


# Formata df de impostos para inseri-los no grupo de despesas formatadas (Impostos sobre Venda / Imposto de Renda)
def formata_impostos_para_dre(df_projecao_impostos, df_orcamentos, casa, mes_selecionado, ano_selecionado, tipo_imposto):
    df_orcamentos_impostos = df_orcamentos[(df_orcamentos['Casa'] == casa) & (df_orcamentos['Classificacao_Contabil_1'] == tipo_imposto)].copy()
    df_orcamentos_impostos = pd.merge(
        df_orcamentos_impostos,
        df_projecao_impostos,
        on=['Ano', 'Mês', 'Categoria'],
        how='left'
    )

    # Cria coluna de Valor Real (apenas para meses passados)
    df_orcamentos_impostos["Valor Real"] = np.where(
        (df_orcamentos_impostos["Mês"] >= datas['mes_atual']) & (df_orcamentos_impostos['Ano'] == datas['ano_atual']),
        None,
        df_orcamentos_impostos["Valor"],                   
    )
    # Cria coluna de Valor Projetado (apenas para meses futuros)
    df_orcamentos_impostos["Valor Projetado"] = np.where(
        (df_orcamentos_impostos["Mês"] >= datas['mes_atual']) & (df_orcamentos_impostos['Ano'] == datas['ano_atual']),
        df_orcamentos_impostos["Valor"],      
        None
    )
    # Cria coluna de Percentual Projetado (apenas para meses futuros)
    df_orcamentos_impostos["Percentual Projetado"] = np.where(
        (df_orcamentos_impostos["Mês"] >= datas['mes_atual']) & (df_orcamentos_impostos['Ano'] == datas['ano_atual']),     
        (df_orcamentos_impostos["Valor"].astype(float) / df_orcamentos_impostos['Orçamento'].astype(float)) * 100,
        None      
    )

    df_impostos_dre = df_orcamentos_impostos[(df_orcamentos_impostos['Mês'] == mes_selecionado) & (df_orcamentos_impostos['Ano'] == ano_selecionado)].copy()
    df_impostos_dre = df_impostos_dre[['Categoria', 'Orçamento', 'Percentual Projetado', 'Valor Projetado', 'Valor Real']]
    
    # Define a mesma ordem dos três impostos 
    if tipo_imposto == 'Impostos sobre Venda':
        ordem = ['PIS / COFINS', 'ICMS', 'ISS']
        df_impostos_dre['Categoria'] = pd.Categorical(df_impostos_dre['Categoria'], categories=ordem, ordered=True)
        df_impostos_dre = df_impostos_dre.sort_values('Categoria')

    df_impostos_dre = calcula_linha_total(df_impostos_dre, 'Categoria', tipo_imposto, 'Valor Projetado', 'Valor Real')
    return df_impostos_dre


############################################ PROJEÇÃO CMV - PRÓXIMOS MESES ############################################

# Soma faturamento de Alimentos, Bebidas e Delivery por casa e mês
def config_faturamento_bruto_zig(df, data_inicio, data_fim, casa):
    df['Valor Bruto'] = df['Valor Bruto'].astype(float)
    df['Data Evento'] = pd.to_datetime(df['Data Evento'], errors='coerce')
    df['Mes_Ano'] = df['Data Evento'].dt.strftime('%Y-%m')
    df = df[
        df['Categoria'].isin(['Alimentos', 'Bebidas', 'Delivery']) &
        (df['Casa'] == casa) &
        (df['Data Evento'] >= data_inicio) &
        (df['Data Evento'] <= data_fim)
    ]

    df = df.groupby(['ID_Casa', 'Casa', 'Mes_Ano', 'Categoria']).agg({
        'Valor Bruto': 'sum',
        'Desconto': 'sum',
        'Valor Liquido': 'sum',
        'Data Evento': 'first'
    }).reset_index()

    faturamento_bruto_alimentos = df[(df['Categoria'] == 'Alimentos')]['Valor Bruto'].sum()
    faturamento_bruto_bebidas = df[(df['Categoria'] == 'Bebidas')]['Valor Bruto'].sum()
    faturamento_delivery = df[(df['Categoria'] == 'Delivery')]['Valor Bruto'].sum()

    return df, faturamento_bruto_alimentos, faturamento_bruto_bebidas, faturamento_delivery


def config_compras(data_inicio, data_fim, loja):
    df1 = GET_INSUMOS_AGRUPADOS_BLUE_ME_POR_CATEG_SEM_PEDIDO()
    df1 = df1.rename(columns={'Loja':'Casa'})
    # Fix 2026-08-11 (mesmo achado/fix de GET_INSUMOS_AGRUPADOS_BLUE_ME_POR_CATEG_COM_
    # PEDIDO_PERIODO_LOJA acima): esta query "sem pedido" também não canonicaliza casa
    # agregada — diferente da versão usada pela aba CMV Real (utils/functions/cmv.py::
    # substituicao_ids), que já resolve isso. Sem esse fix, itens "sem pedido" da empresa
    # 131/178 (Blue Note SP Novo/Sala 2) ficavam fora da soma de Compras do Forecast.
    df1['Casa'] = df1['Casa'].replace({
        'Blue Note SP (Novo)': 'Blue Note - São Paulo',
        'Blue Note SP (Sala 2)': 'Blue Note - São Paulo',
    })
    df1['Primeiro_Dia_Mes'] = pd.to_datetime(df1['Primeiro_Dia_Mes'], errors='coerce')
    df1['Mes_Ano'] = df1['Primeiro_Dia_Mes'].dt.strftime('%Y-%m')
    # Fix 2026-08-11 (correção do fix acima): a canonicalização de nome por si só não
    # reagrupa — Blue Note passou a ter 2 linhas com o mesmo 'Casa' (uma por empresa
    # bruta, 110 e 131), e o merge abaixo (chave Casa+Mes_Ano+Primeiro_Dia_Mes) casava
    # CADA uma delas contra a mesma linha já agregada de df2, duplicando (contando 2x)
    # o valor "com pedido" — achado real: CMV Real subiu de 24,37% (faltando compras) p/
    # 61,91% (compras "com pedido" contadas em dobro) no mesmo período. Agrega df1 por
    # Casa/Mes antes do merge, mesmo problema resolvido no lado SQL de df2 pelo GROUP BY
    # já usar as colunas canonicalizadas.
    df1 = df1.groupby(['Casa', 'Primeiro_Dia_Mes', 'Mes_Ano'], as_index=False).agg({
        'BlueMe_Sem_Pedido_Valor': 'sum',
        'BlueMe_Sem_Pedido_Alimentos': 'sum',
        'BlueMe_Sem_Pedido_Bebidas': 'sum',
        'BlueMe_Sem_Pedido_Descart_Hig_Limp': 'sum',
        'BlueMe_Sem_Pedido_Outros': 'sum',
    })

    df_aut_blue_me_com_pedido = GET_AUT_BLUE_ME_COM_PEDIDO()
    df2 = GET_INSUMOS_AGRUPADOS_BLUE_ME_POR_CATEG_COM_PEDIDO_PERIODO_LOJA()
    df2['Primeiro_Dia_Mes'] = pd.to_datetime(df2['Primeiro_Dia_Mes'], errors='coerce')
    df2['Mes_Ano'] = df2['Primeiro_Dia_Mes'].dt.strftime('%Y-%m')

    df_compras = pd.merge(df2, df1, on=['Casa', 'Mes_Ano', 'Primeiro_Dia_Mes'], how='outer')

    df_compras = df_compras[
        (df_compras['Casa'] == loja) &
        (df_compras['Primeiro_Dia_Mes'] >= data_inicio) &
        (df_compras['Primeiro_Dia_Mes'] <= data_fim)
    ].copy()

    df_compras = df_compras.groupby(['Casa', 'Mes_Ano']).agg(
        {'BlueMe_Sem_Pedido_Alimentos': 'sum', 
        'BlueMe_Sem_Pedido_Bebidas': 'sum', 
        'Valor_Liq_Alimentos': 'sum', 
        'Valor_Liq_Bebidas': 'sum',
        'BlueMe_Sem_Pedido_Descart_Hig_Limp': 'sum'}).reset_index()

    Compras_Alimentos = df_compras['BlueMe_Sem_Pedido_Alimentos'].sum() + df_compras['Valor_Liq_Alimentos'].sum()
    Compras_Bebidas = df_compras['BlueMe_Sem_Pedido_Bebidas'].sum() + df_compras['Valor_Liq_Bebidas'].sum()

    Compras_Alimentos = float(Compras_Alimentos)
    Compras_Bebidas = float(Compras_Bebidas)

    df_compras['Compras Alimentos'] = df_compras['Valor_Liq_Alimentos'] + df_compras['BlueMe_Sem_Pedido_Alimentos']
    df_compras['Compras Bebidas'] = df_compras['Valor_Liq_Bebidas'] + df_compras['BlueMe_Sem_Pedido_Bebidas']
    df_compras['Compras Embalagens'] = df_compras['BlueMe_Sem_Pedido_Descart_Hig_Limp']
    df_compras = df_compras.rename(columns={'Valor_Liq_Alimentos': 'BlueMe c/ Pedido Alim.', 'Valor_Liq_Bebidas': 'BlueMe c/ Pedido Bebidas', 'BlueMe_Sem_Pedido_Alimentos': 'BlueMe s/ Pedido Alim.', 'BlueMe_Sem_Pedido_Bebidas': 'BlueMe s/ Pedido Bebidas'})

    df_compras = df_compras[['Casa', 'Mes_Ano', 'BlueMe c/ Pedido Alim.', 'BlueMe s/ Pedido Alim.', 'Compras Alimentos', 'BlueMe c/ Pedido Bebidas', 'BlueMe s/ Pedido Bebidas', 'Compras Bebidas', 'Compras Embalagens']]

    return df_compras, df_aut_blue_me_com_pedido, Compras_Alimentos, Compras_Bebidas


def processar_transferencias(df, casa_col, loja, data_inicio, data_fim):
    # Filtrar pelo nome da loja e pelo intervalo de datas
    df['Data_Transferencia'] = pd.to_datetime(df['Data_Transferencia'], errors='coerce')
    df['Mes_Ano'] = df['Data_Transferencia'].dt.strftime('%Y-%m')
    df = df[
        (df[casa_col] == loja) &
        (df['Data_Transferencia'] >= data_inicio) &
        (df['Data_Transferencia'] <= data_fim)
    ].copy()
    
    # Agrupar por casa e categoria, somando os valores
    df_grouped = df.groupby([casa_col, 'Categoria', 'Mes_Ano']).agg({
        'Valor_Transferencia': 'sum'
    }).reset_index()
    
    # Ajustar categoria para formato capitalizado
    df_grouped['Categoria'] = df_grouped['Categoria'].str.capitalize()
    
    # Pivotar para transformar categorias em colunas
    df_pivot = df_grouped.pivot_table(
        index=[casa_col, 'Mes_Ano'],
        columns='Categoria',
        values='Valor_Transferencia',
        fill_value=0
    ).reset_index()
    
    # Renomear colunas para refletir o tipo de operação
    operacao = 'Entrada' if casa_col == 'Casa_Entrada' else 'Saída'
    df_pivot.columns = [
        'Loja' if col == casa_col else col if col == 'Mes_Ano' else f'{operacao} {col}'
        for col in df_pivot.columns
    ]    
    return df_pivot

def config_transferencias_gastos(data_inicio, data_fim, loja):
    df_transf_estoque = GET_TRANSF_ESTOQUE()
    df_entradas_pivot = processar_transferencias(df_transf_estoque, 'Casa_Entrada', loja, data_inicio, data_fim)
    df_saidas_pivot = processar_transferencias(df_transf_estoque, 'Casa_Saida', loja, data_inicio, data_fim)
    
    df_perdas_e_consumo = GET_PERDAS_E_CONSUMO_AGRUPADOS()
    df_perdas_e_consumo['Primeiro_Dia_Mes'] = pd.to_datetime(df_perdas_e_consumo['Primeiro_Dia_Mes'], errors='coerce')
    df_perdas_e_consumo['Mes_Ano'] = df_perdas_e_consumo['Primeiro_Dia_Mes'].dt.strftime('%Y-%m')
    df_perdas_e_consumo = df_perdas_e_consumo[
        (df_perdas_e_consumo['Loja'] == loja) &
        (df_perdas_e_consumo['Primeiro_Dia_Mes'] >= data_inicio) &
        (df_perdas_e_consumo['Primeiro_Dia_Mes'] <= data_fim)
    ].copy()
    df_perdas_e_consumo = df_perdas_e_consumo.fillna(0)

    df_transf_e_gastos = pd.merge(df_entradas_pivot, df_saidas_pivot, on=['Loja', 'Mes_Ano'], how='outer')
    df_transf_e_gastos = pd.merge(df_transf_e_gastos, df_perdas_e_consumo, on=['Loja', 'Mes_Ano'], how='outer')
    df_transf_e_gastos = df_transf_e_gastos.rename(columns={
        'ID_Loja': 'ID_Casa',
        'Loja': 'Casa',
        'Consumo_Interno': 'Consumo Interno',
        'Quebras_e_Perdas': 'Quebras e Perdas'
    })
    cols = ['ID_Casa', 'Casa', 'Mes_Ano', 'Entrada Alimentos', 'Entrada Bebidas', 'Saída Alimentos', 'Saída Bebidas', 'Consumo Interno', 'Quebras e Perdas']
    for col in cols:
        if col not in df_transf_e_gastos.columns:
            df_transf_e_gastos[col] = 0

    df_transf_e_gastos = df_transf_e_gastos[cols]

    # Conversão para float para evitar erros de tipo
    # saida_alimentos = float(df_saidas_pivot['Saída Alimentos'].iloc[0]) if not df_saidas_pivot.empty and 'Saída Alimentos' in df_saidas_pivot.columns else 0.0
    # saida_bebidas = float(df_saidas_pivot['Saída Bebidas'].iloc[0]) if not df_saidas_pivot.empty and 'Saída Bebidas' in df_saidas_pivot.columns else 0.0
    # entrada_alimentos = float(df_entradas_pivot['Entrada Alimentos'].iloc[0]) if not df_entradas_pivot.empty and 'Entrada Alimentos' in df_entradas_pivot.columns else 0.0
    # entrada_bebidas = float(df_entradas_pivot['Entrada Bebidas'].iloc[0]) if not df_entradas_pivot.empty and 'Entrada Bebidas' in df_entradas_pivot.columns else 0.0
    # consumo_interno = float(df_transf_e_gastos['Consumo Interno'].iloc[0]) if not df_perdas_e_consumo.empty and 'Consumo Interno' in df_transf_e_gastos.columns else 0.0
    # quebras_e_perdas = float(df_transf_e_gastos['Quebras e Perdas'].iloc[0]) if not df_perdas_e_consumo.empty and 'Quebras e Perdas' in df_transf_e_gastos.columns else 0.0

    return df_transf_e_gastos #, saida_alimentos, saida_bebidas, entrada_alimentos, entrada_bebidas, consumo_interno, quebras_e_perdas


def config_valoracao_estoque_ou_producao(tipo, data_inicio, data_fim, loja):
    # Pega os dados
    if tipo == 'estoque':
        df_valoracao = GET_VALORACAO_ESTOQUE(data_inicio, data_fim)
        col_data = 'DATA_CONTAGEM'
        col_valor = 'Valor_em_Estoque'
    if tipo == 'producao':
        df_valoracao = GET_VALORACAO_PRODUCAO(data_inicio, data_fim)
        col_data = 'Data_Contagem'
        col_valor = 'Valor_Total'
    
    df_valoracao = df_valoracao[
       (df_valoracao['Loja'] == loja) &
       (df_valoracao['Categoria'].isin(['ALIMENTOS', 'BEBIDAS']))
    ].copy()

    # Fix 2026-08-11 revertido em 2026-09-02 (decisão do usuário, mesmo racional de
    # utils/functions/cmv.py::config_valoracao_estoque): GET_VALORACAO_ESTOQUE traz a
    # linha crua sem dedupe. Até aqui, linhas com mesma loja+insumo+quantidade+valor(+
    # DATA_CONTAGEM) eram tratadas como reenvio duplicado da mesma contagem e
    # descartadas automaticamente antes de somar por mês. Com múltiplas contagens por
    # casa em pontos de estoque diferentes (BlueMe), essa combinação pode ser 2
    # contagens físicas legítimas — decidir sozinho ficou arriscado demais, agora soma
    # tudo como está. Se for de fato reenvio, a correção é excluir a contagem errada em
    # T_CONTAGEM_INSUMOS (ver flag no script Mini_Gabu/BlueMe Dashboard) e reapurar.

    # Fix 2026-08-10: contagens recentes (lag de consolidação — mesmo fenômeno já
    # documentado nos fixes de GET_VALORACAO_PRODUCAO/DRE_AUT_INSUMOS_PRODUCAO) chegam
    # com o horário exato de submissão em vez de meia-noite, e a query de estoque
    # (GET_VALORACAO_ESTOQUE) mistura DATETIME cru com string formatada (DATE_FORMAT)
    # dependendo se a contagem é agrupada — pd.to_datetime com formato misto pode
    # travar no primeiro formato inferido e descartar (NaT) todo o resto silenciosamente.
    # Isso fazia o fechamento do mês mais recente (ex.: produção com dois horários
    # distintos no mesmo dia, um por categoria) não casar com o "scaffold" de meses
    # abaixo — Variação_Mensal do mês virava 0 - valor_anterior em vez de atual -
    # anterior, distorcendo o CMV Real da aba Forecast (achado real: Bar Brahma - Granja,
    # jul/2026 — CMV saía -35,64% em vez de -30,36%, mesmo valor já validado no CMV
    # Real/Download DRE). Normaliza para os 10 primeiros caracteres (parte de data, sem
    # horário) antes de parsear e já trunca para o 1º dia do mês — garante o match com o
    # scaffold independente do horário exato de submissão ou do formato retornado.
    df_valoracao[col_data] = pd.to_datetime(
        df_valoracao[col_data].astype(str).str[:10], errors='coerce'
    ).dt.to_period('M').dt.to_timestamp()

    # Agrupar por mês, loja e categoria
    df_valoracao = (
        df_valoracao
        .groupby(['Loja', 'Categoria', col_data], as_index=False)
        .agg({col_valor: 'sum'})
    )

    # Criar todas as datas do período
    todas_datas = pd.date_range(start=data_inicio, end=data_fim, freq='MS')

    # Todas combinações de Loja, Categoria e DATA_CONTAGEM
    lojas_categorias = df_valoracao[['Loja', 'Categoria']].drop_duplicates()
    todas_combinacoes = (
        lojas_categorias
        .merge(pd.DataFrame({col_data: todas_datas}), how='cross')
    )

    # Merge com o dataframe real
    df_valoracao = todas_combinacoes.merge(
        df_valoracao,
        on=['Loja', 'Categoria', col_data],
        how='left'
    )
    
    # Preencher valores ausentes com 0
    df_valoracao[col_valor] = df_valoracao[col_valor].fillna(0)
    
    # Calcular variação mensal
    df_valoracao['Variação_Mensal'] = (
        df_valoracao
        .groupby(['Loja', 'Categoria'])[col_valor]
        .diff()
        .fillna(0)
    )
    
    # Coluna DATA_MES_ANTERIOR
    df_valoracao['DATA_MES_ANTERIOR'] = df_valoracao[col_data] - pd.DateOffset(months=1)
    df_valoracao['Mes_Ano'] = df_valoracao['DATA_MES_ANTERIOR'].dt.strftime("%Y-%m")
    df_valoracao = df_valoracao.rename(columns={'Loja': 'Casa'})

    return df_valoracao


def config_faturamento_eventos(data_inicio, data_fim, loja, faturamento_bruto_alimentos, faturamento_bruto_bebidas):
    df_eventos = GET_EVENTOS_CMV(data_inicio=data_inicio, data_fim=data_fim)
    
    df_eventos = df_eventos[df_eventos['Loja'] == loja]
    df_eventos['Valor_AB'] = df_eventos['Valor_AB'].astype(float)
    df_eventos['Data'] = pd.to_datetime(df_eventos['Data'], errors='coerce')
    df_eventos['Mes_Ano'] = df_eventos['Data'].dt.strftime('%Y-%m')
    df_eventos = df_eventos.rename(columns={'ID_Loja': 'ID_Casa', 'Loja': 'Casa'})

    return df_eventos


def merge_e_calculo_para_cmv(df_faturamento_zig, df_compras, df_valoracao_estoque, df_transf_e_gastos, df_valoracao_producao, df_faturamento_eventos, df_ajustes_manuais, casa, ano_selecionado):
    # Faturamento Bruto (alimentos + bebidas + delivery) mensal
    df_faturamento_zig_geral = df_faturamento_zig.copy()
    df_faturamento_zig_geral = df_faturamento_zig_geral.groupby(['ID_Casa', 'Casa', 'Mes_Ano'], as_index=False)['Valor Bruto'].sum()
    df_faturamento_zig_geral = df_faturamento_zig_geral.rename(columns={'Valor Bruto':'Faturamento Bruto'})

    # Compras (alimentos + bebidas) mensais
    df_compras_geral = df_compras.copy()
    df_compras_geral['Compras Geral'] = df_compras_geral['Compras Alimentos'] + df_compras_geral['Compras Bebidas'] + df_compras_geral['Compras Embalagens']
    df_compras_geral = df_compras_geral[['Casa', 'Mes_Ano', 'Compras Geral']]

    # Valoração estoque (alimentos + bebidas) mensal
    df_valoracao_estoque_geral = df_valoracao_estoque.copy()
    df_valoracao_estoque_geral = df_valoracao_estoque_geral.groupby(['Casa', 'Mes_Ano'], as_index=False)['Variação_Mensal'].sum()

    # Transferências e gastos (entrada alimentos + bebidas, saidas alimentos + bebidas, consumo interno, quebras e perdas)
    df_transf_e_gastos_geral = df_transf_e_gastos.copy()
    df_transf_e_gastos_geral['Entradas Geral'] = df_transf_e_gastos_geral['Entrada Alimentos'] + df_transf_e_gastos_geral['Entrada Bebidas']
    df_transf_e_gastos_geral['Saídas Geral'] = df_transf_e_gastos_geral['Saída Alimentos'] + df_transf_e_gastos_geral['Saída Bebidas']
    df_transf_e_gastos_geral = df_transf_e_gastos_geral[['ID_Casa', 'Casa', 'Mes_Ano', 'Entradas Geral', 'Saídas Geral', 'Consumo Interno', 'Quebras e Perdas']]

    # Valoração produção (alimentos + bebidas) mensal
    df_valoracao_producao_geral = df_valoracao_producao.copy()
    df_valoracao_producao_geral = df_valoracao_producao_geral.groupby(['Casa', 'Mes_Ano'], as_index=False)['Variação_Mensal'].sum()

    # Faturamento eventos (A&B) mensal
    df_faturamento_eventos_geral = df_faturamento_eventos.copy()
    df_faturamento_eventos_geral = df_faturamento_eventos_geral.groupby(['ID_Casa', 'Casa', 'Mes_Ano'], as_index=False)['Valor_AB'].sum()

    # Merge passo a passo para cálculo do cmv
    df_cmv = (
        df_compras_geral
            .merge(df_valoracao_estoque_geral, on=['Casa', 'Mes_Ano'], how='left')
            .merge(df_transf_e_gastos_geral, on=['Casa', 'Mes_Ano'], how='left')
            .merge(df_valoracao_producao_geral, on=['Casa', 'Mes_Ano'], how='left')
    ).fillna(0)
    df_cmv = df_cmv.rename(columns={'Variação_Mensal_x':'Variacao_Estoque', 'Variação_Mensal_y':'Variacao_Producao'})

    # Merge com ajustes manuais de CMV 
    df_cmv['Mes_Ano_copia'] = pd.to_datetime(df_cmv['Mes_Ano'])
    df_cmv['Ano'] = df_cmv['Mes_Ano_copia'].dt.year
    df_cmv['Mês'] = df_cmv['Mes_Ano_copia'].dt.month

    df_ajustes_cmv = df_ajustes_manuais[
        (df_ajustes_manuais['Casa'] == casa) &
        (df_ajustes_manuais['Ano'] == ano_selecionado) &
        (df_ajustes_manuais['Classificacao_Contabil_1'] == 'Custo Mercadoria Vendida')
    ].copy()
    
    df_ajustes_cmv = df_ajustes_cmv.groupby(['ID_Casa', 'Casa', 'Mês', 'Ano', 'Classificacao_Contabil_1'], as_index=False)['Valor Ajuste'].sum()

    if not df_ajustes_cmv.empty:
        df_cmv_com_ajustes = pd.merge(
            df_cmv,
            df_ajustes_cmv,
            on=['Casa', 'Ano', 'Mês'],
            how='left'
        )
        mask = ~df_cmv_com_ajustes['Valor Ajuste'].isna()
        df_cmv_com_ajustes.loc[mask, 'Compras Geral'] -= (df_cmv_com_ajustes.loc[mask, 'Valor Ajuste']) # Soma valores negativos e subtrai positivos
    else:
        df_cmv_com_ajustes = df_cmv.copy()

    # Faturamento geral (bruto + eventos)
    df_merge_faturamento = df_faturamento_zig_geral.merge(
        df_faturamento_eventos_geral,
        on=['Casa', 'Mes_Ano'],
        how='left'
    ).fillna(0)
    df_merge_faturamento['Faturamento_Geral'] = df_merge_faturamento['Faturamento Bruto'] + df_merge_faturamento['Valor_AB']
    
    df_merge_cmv = pd.merge(
        df_cmv_com_ajustes[['Casa', 'Mes_Ano', 'Compras Geral', 'Variacao_Estoque', 'Entradas Geral', 'Saídas Geral', 'Consumo Interno', 'Quebras e Perdas', 'Variacao_Producao']],
        df_merge_faturamento[['Casa', 'Mes_Ano', 'Faturamento_Geral']],
        on=['Casa', 'Mes_Ano'],
        how='right'
    )

    # Cálculo do cmv e porcentagem cmv
    df_calculo_cmv = df_merge_cmv.copy()
    df_calculo_cmv['Compras Geral'] = df_calculo_cmv['Compras Geral'].astype(float)
    df_calculo_cmv['Variacao_Estoque'] = df_calculo_cmv['Variacao_Estoque'].astype(float)
    df_calculo_cmv['Entradas Geral'] = df_calculo_cmv['Entradas Geral'].astype(float)
    df_calculo_cmv['Saídas Geral'] = df_calculo_cmv['Saídas Geral'].astype(float)
    df_calculo_cmv['Consumo Interno'] = df_calculo_cmv['Consumo Interno'].astype(float)
    df_calculo_cmv['Variacao_Producao'] = df_calculo_cmv['Variacao_Producao'].astype(float)

    df_calculo_cmv['CMV Real'] = df_calculo_cmv['Compras Geral'] - df_calculo_cmv['Variacao_Estoque'] + df_calculo_cmv['Entradas Geral'] - df_calculo_cmv['Saídas Geral'] - df_calculo_cmv['Consumo Interno'] - df_calculo_cmv['Variacao_Producao']
    df_calculo_cmv['CMV Real Percentual'] = (df_calculo_cmv['CMV Real'] / df_calculo_cmv['Faturamento_Geral']) * 100
    # st.write('cmv', df_calculo_cmv)
    df_calculo_cmv = df_calculo_cmv[['Casa', 'Mes_Ano', 'Faturamento_Geral', 'CMV Real', 'CMV Real Percentual']]

    return df_calculo_cmv


# Utiliza o df de faturamento projetado criado anteriormente (para projetar o cmv para os prox meses)
def calcula_cmv_proximos_meses(df_faturamento_meses_futuros, df_calculo_cmv, ano_atual, mes_atual):
    df_resgata_faturamento_meses_futuros = df_faturamento_meses_futuros[
        (df_faturamento_meses_futuros['Ano'] >= ano_atual - 1) &
        (df_faturamento_meses_futuros['Categoria'].isin(['Alimentos', 'Bebidas', 'Delivery', 'Eventos A&B']))
    ].copy()
    
    # Resgata faturamentos projetados por mês
    df_resgata_faturamento_meses_futuros['Valor Bruto'] = pd.to_numeric(df_resgata_faturamento_meses_futuros['Valor Bruto'], errors='coerce')
    df_resgata_faturamento_meses_futuros['Valor Projetado'] = pd.to_numeric(df_resgata_faturamento_meses_futuros['Valor Projetado'], errors='coerce')
    df_resgata_faturamento_meses_futuros = df_resgata_faturamento_meses_futuros.groupby(['Ano', 'Mês'], as_index=False)[['Valor Bruto', 'Valor Projetado']].sum()
    df_resgata_faturamento_meses_futuros['Mês'] = df_resgata_faturamento_meses_futuros['Mês'].astype(int)
    df_resgata_faturamento_meses_futuros['Mes_Ano'] = (df_resgata_faturamento_meses_futuros['Ano'].astype(int).astype(str) + '-' + df_resgata_faturamento_meses_futuros['Mês'].astype(int).astype(str).str.zfill(2))    
    df_merge_meses_anteriores_seguintes = pd.merge(
        df_calculo_cmv,
        df_resgata_faturamento_meses_futuros[['Ano', 'Mês', 'Mes_Ano', 'Valor Projetado']],
        on='Mes_Ano',
        how='right'
    )
    df_merge_meses_anteriores_seguintes['Data'] = pd.to_datetime(df_merge_meses_anteriores_seguintes['Mes_Ano'], format='%Y-%m')
    
    # Cria coluna para CMV projetado de cada mês
    df_merge_meses_anteriores_seguintes['CMV Percentual Projetado'] = None
    df_merge_meses_anteriores_seguintes['CMV Projetado'] = None
    
    for mes_ano in df_merge_meses_anteriores_seguintes['Mes_Ano'].unique():
        df_mes_ano = df_merge_meses_anteriores_seguintes[df_merge_meses_anteriores_seguintes['Mes_Ano'] == mes_ano]
        data = df_mes_ano['Data'].iloc[0]

        # pega histórico dos dois meses atrás
        dois_meses_atras = data - pd.DateOffset(months=2)
        
        historico = df_merge_meses_anteriores_seguintes[
            (df_merge_meses_anteriores_seguintes['Data'] >= dois_meses_atras) &
            (df_merge_meses_anteriores_seguintes['Data'] < data)
        ].copy()
        
        # Faz Projecao = (CMV1 + CMV2) / (Faturamento_Geral1 + Faturamento_Geral2)
        # Define colunas auxiliares conforme o mês
        historico["CMV_Usado"] = np.where(
            (historico["Mês"] >= mes_atual) & (historico['Ano'] == ano_atual),
            historico["CMV Projetado"],       # usa o projetado se mês >= atual
            historico["CMV Real"]                  # senão usa o real
        )

        historico["Faturamento_Usado"] = np.where(
           (historico["Mês"] >= mes_atual) & (historico['Ano'] == ano_atual),
            historico["Valor Projetado"],          # usa o projetado se mês >= atual
            historico["Faturamento_Geral"]         # senão usa o real
        )
        
        valores_para_soma_cmvs = historico['CMV_Usado'].fillna(historico['CMV Projetado']).astype(float)
        valores_para_soma_faturamento = historico['Faturamento_Usado'].fillna(historico['Valor Projetado']).astype(float)

        soma_cmvs = valores_para_soma_cmvs.sum()
        soma_faturamentos = valores_para_soma_faturamento.sum()
        
        if soma_faturamentos and not pd.isna(soma_faturamentos) and soma_faturamentos != 0:
            cmv_projetado = (soma_cmvs / soma_faturamentos) * 100
        else:
            cmv_projetado = 0 

        # Atribui o valor à coluna correta
        df_merge_meses_anteriores_seguintes.loc[
            df_merge_meses_anteriores_seguintes['Mes_Ano'] == mes_ano, 
            'CMV Percentual Projetado'
        ] = cmv_projetado

        # Define valor de CMV Projetado em Reais
        df_merge_meses_anteriores_seguintes['CMV Projetado'] = (df_merge_meses_anteriores_seguintes['CMV Percentual Projetado'] / 100) * df_merge_meses_anteriores_seguintes['Valor Projetado']

    return df_merge_meses_anteriores_seguintes


############################################ PROJEÇÃO DESPESAS - PRÓXIMOS MESES ############################################

def merge_despesas_complexas(df_tabela_primaria, df_tabela_secundaria, df_tabela_terciaria, df_tabela_quaternaria, casa, class_cont):
    df_tabela_secundaria_filtrada = df_tabela_secundaria[df_tabela_secundaria['Casa'] == casa].copy()

    if class_cont in ['Custos Artístico Geral', 'Marketing', 'Informática e TI', 'Mão de Obra - Benefícios']: # merge com Descontos
        df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada.groupby(['Casa', 'Mês', 'Ano', 'Centro de Custo'], as_index=False)['Aloca no Centro de Custo'].sum()
        df_tabela_resultante = pd.merge(
            df_tabela_primaria, df_tabela_secundaria_filtrada,
            left_on=['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2'],
            right_on=['Casa', 'Mês', 'Ano', 'Centro de Custo'],
            how='outer'
        )

        if class_cont == 'Custos Artístico Geral': categorias_consideradas_descontos = ['Alimentação e Transporte']
        elif class_cont == 'Marketing': categorias_consideradas_descontos = ['Eventos de Marketing', 'Produção Gráfica e Material Institucional', 'Ferramentas de Marketing']
        elif class_cont == 'Informática e TI': categorias_consideradas_descontos = ['Sistemas Gerais - Operacionais']
        elif class_cont == 'Mão de Obra - Benefícios': categorias_consideradas_descontos = ['  -  Alimentação Funcionário']

        condicao = (df_tabela_resultante['Centro de Custo'].isin(categorias_consideradas_descontos) & df_tabela_resultante['Classificacao_Contabil_2'].isna())
        df_tabela_resultante.loc[condicao, 'Classificacao_Contabil_2'] = df_tabela_resultante['Centro de Custo']
        df_tabela_resultante['Custo Real'] = df_tabela_resultante['Custo Real'].fillna(0)
        df_tabela_resultante['Aloca no Centro de Custo'] = df_tabela_resultante['Aloca no Centro de Custo'].fillna(0)

        condicao = df_tabela_resultante['Classificacao_Contabil_2'].isin(['Alimentação e Transporte', 'Eventos de Marketing', 'Produção Gráfica e Material Institucional', 'Ferramentas de Marketing', 'Sistemas Gerais - Operacionais', '  -  Alimentação Funcionário'])
        df_tabela_resultante.loc[condicao, 'Custo Real'] = df_tabela_resultante['Custo Real'] + df_tabela_resultante['Aloca no Centro de Custo']
        df_tabela_resultante = df_tabela_resultante.drop(columns=['Centro de Custo', 'Aloca no Centro de Custo'])
        df_tabela_resultante = df_tabela_resultante.dropna(subset=['Classificacao_Contabil_2'])

        if class_cont == 'Mão de Obra - Benefícios': # Alimentação Funcionário envolve CMV e Cartão Black
            # Consumo Interno - CMV
            if casa not in ['Blue Note - São Paulo']:
                df_tabela_terciaria['Classificacao_Contabil_2'] = '  -  Alimentação Funcionário'
                df_tabela_terciaria = df_tabela_terciaria[['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2', 'Consumo Interno']]
                df_tabela_terciaria = df_tabela_terciaria.rename(columns={'Consumo Interno': 'Custo Real'})
                df_tabela_resultante = pd.concat([df_tabela_resultante, df_tabela_terciaria])
                df_tabela_resultante = df_tabela_resultante.groupby(['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2'], as_index=False)['Custo Real'].sum()

            # Consumo - Cartão Black
            if casa not in ['Arcos', 'Blue Note - São Paulo', 'Love Cabaret']: # Era pra adicionar 'Ultra Evil Premium Ltda '
                df_tabela_quaternaria_filtrada = df_tabela_quaternaria[df_tabela_quaternaria['Casa'] == casa].copy()
                df_tabela_quaternaria_filtrada = df_tabela_quaternaria_filtrada.groupby(['Casa', 'Mês', 'Ano'], as_index=False)['Valor Cartão Black'].sum()
                df_tabela_resultante = pd.merge(
                    df_tabela_resultante, df_tabela_quaternaria_filtrada[['Casa', 'Mês', 'Ano', 'Valor Cartão Black']], 
                    on=['Casa', 'Mês', 'Ano'], 
                    how='left'
                )
                df_tabela_resultante['Custo Real'] = df_tabela_resultante['Custo Real'].fillna(0)
                df_tabela_resultante['Valor Cartão Black'] = df_tabela_resultante['Valor Cartão Black'].fillna(0)
                condicao = df_tabela_resultante['Classificacao_Contabil_2'] == '  -  Alimentação Funcionário'
                df_tabela_resultante.loc[condicao, 'Custo Real'] = df_tabela_resultante['Custo Real'] + df_tabela_resultante['Valor Cartão Black']
                df_tabela_resultante = df_tabela_resultante.drop(columns=['Valor Cartão Black'])

    elif class_cont in ['Gorjeta', 'Mão de Obra - Salários']: # merge com folha de pagamento
        df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada.groupby(['Casa', 'Mês', 'Ano'], as_index=False)['Valor'].sum()
        df_tabela_resultante = pd.merge(
            df_tabela_primaria,
            df_tabela_secundaria_filtrada,
            on=['Casa', 'Mês', 'Ano'],
            how='left'
        ).fillna(0)

        if class_cont == 'Gorjeta':
            df_tabela_resultante['Custo Real'] = df_tabela_resultante['Custo Real'] + df_tabela_resultante['Valor']
        elif class_cont == 'Mão de Obra - Salários':
            df_tabela_resultante['Custo Real'] = df_tabela_resultante['Custo Real'] - df_tabela_resultante['Valor']
        df_tabela_resultante = df_tabela_resultante.drop(columns=['Valor'])

    elif class_cont == 'Patrocínio': # (+) Receitas de Patrocínio (envolve Receitas Extraordinárias)
        df_tabela_secundaria_filtrada['Mês'] = df_tabela_secundaria_filtrada['Recebimento_Parcela'].dt.month
        df_tabela_secundaria_filtrada['Ano'] = df_tabela_secundaria_filtrada['Recebimento_Parcela'].dt.year
        df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada.groupby(['Casa', 'Mês', 'Ano'], as_index=False)['Valor Bruto'].sum()
        df_tabela_secundaria_filtrada['Classificacao_Contabil_2'] = '(+) Receitas de Patrocínio'
        df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada[['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2', 'Valor Bruto']]
        df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada.rename(columns={'Valor Bruto': 'Custo Real'})
        df_tabela_resultante = pd.concat([df_tabela_primaria, df_tabela_secundaria_filtrada])
        df_tabela_resultante = df_tabela_resultante.groupby(['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2'], as_index=False)['Custo Real'].sum()
        if casa == 'Blue Note - São Paulo': # (-) Despesas de Patrocínio (envolve Descontos)
            df_tabela_terciaria_filtrada = df_tabela_terciaria[df_tabela_terciaria['Casa'] == 'Blue Note - São Paulo'].copy()
            df_tabela_terciaria_filtrada = df_tabela_terciaria_filtrada.groupby(['Casa', 'Mês', 'Ano'], as_index=False)['Permanece no Desconto'].sum()
            df_tabela_resultante = pd.merge(df_tabela_resultante, df_tabela_terciaria_filtrada, on=['Casa', 'Mês', 'Ano'], how='left').fillna(0)
            df_tabela_resultante.loc[df_tabela_resultante['Classificacao_Contabil_2'] == '(-) Despesas de Patrocínio', 'Custo Real'] += df_tabela_resultante['Permanece no Desconto']
            df_tabela_resultante = df_tabela_resultante.drop(columns=['Permanece no Desconto'])

    elif class_cont == 'Despesas Financeiras': 
        # merge com outra class. cont. da aut blueme sem pedido
        if casa in ['Bar Brahma - Centro', 'Bar Léo - Centro']: # Aluguel de Imóveis
            df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada[df_tabela_secundaria_filtrada['Classificacao_Contabil_2'] == 'Aluguel de Imoveis'].copy()
            df_tabela_secundaria_filtrada['Data_Competencia'] = pd.to_datetime(df_tabela_secundaria_filtrada['Data_Competencia'], errors='coerce')
            df_tabela_secundaria_filtrada['Mês'] = df_tabela_secundaria_filtrada['Data_Competencia'].dt.month
            df_tabela_secundaria_filtrada['Ano'] = df_tabela_secundaria_filtrada['Data_Competencia'].dt.year
            df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada.groupby(['Casa', 'Mês', 'Ano'], as_index=False)[['Valor_Pagamento', 'Valor_Liquido']].sum()
            df_tabela_resultante = pd.merge(df_tabela_primaria, df_tabela_secundaria_filtrada[['Casa', 'Mês', 'Ano', 'Valor_Pagamento', 'Valor_Liquido']],  on=['Casa', 'Mês', 'Ano'], how='left')
            df_tabela_resultante['Custo Real'] = (df_tabela_resultante['Custo Real'] - df_tabela_resultante['Valor_Pagamento'] + df_tabela_resultante['Valor_Liquido_y'])
            # df_tabela_resultante.loc[df_tabela_resultante['Custo Real'] < 0, 'Custo Real'] *= (-1)
        # merge com Receitas Extr. Blue Note - Bilheteria-Rebate
        elif casa in ['Blue Note - São Paulo']:
            df_tabela_secundaria_filtrada['Mês'] = df_tabela_secundaria_filtrada['Recebimento_Parcela'].dt.month
            df_tabela_secundaria_filtrada['Ano'] = df_tabela_secundaria_filtrada['Recebimento_Parcela'].dt.year
            df_tabela_secundaria_filtrada = df_tabela_secundaria_filtrada.groupby(['Casa', 'Classificacao', 'Mês', 'Ano'], as_index=False)[['Valor Bruto']].sum()
            df_bilheteria_rebate = df_tabela_secundaria_filtrada[df_tabela_secundaria_filtrada['Classificacao'] == 'Bilheteria-Rebate'].copy()
            df_cashback = df_tabela_secundaria_filtrada[df_tabela_secundaria_filtrada['Classificacao'] == 'Cashback'].copy()
            df_tabela_resultante = (df_tabela_primaria.merge(df_bilheteria_rebate, on=['Casa', 'Ano', 'Mês'], how='left').merge(df_cashback, on=['Casa', 'Ano', 'Mês'], how='left')).fillna(0)
            df_tabela_resultante['Custo Real'] = ((df_tabela_resultante['Custo Real'] * (-1)) + df_tabela_resultante['Valor Bruto_x'] + df_tabela_resultante['Valor Bruto_y']) * (-1) # Para ficar positivo no layout final, caso seja
        else: df_tabela_resultante = df_tabela_primaria.copy() # Outras casas não precisam do merge   

    return df_tabela_resultante


def prepara_dados_custos_mensais(df_custos_gerais, df_faturamento_meses_futuros, casa, class_cont, df_orcamentos, df_valor_fee_gestao, df_aut_blue_me_com_pedido=None, df_tabela_secundaria=None, df_tabela_terciaria=None, df_tabela_quaternaria=None):
    # Filtra pela casa
    df_custos_filtrado = df_custos_gerais[df_custos_gerais['Casa'] == casa ].copy()

    # Filtra por class. cont. 1
    if class_cont == 'Custos Artístico Geral': # Realoca MDO de PJ para Artístico
        if casa == 'Ultra Evil Premium Ltda ':
            df_custos_filtrado = df_custos_filtrado[
                ((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
                (df_custos_filtrado['Classificacao_Contabil_2'] != 'MDO Terceirizada - Artístico')) |
                (df_custos_filtrado['Fornecedor'] == 'JEFERSON LUIS DE GODOI ')
            ].copy()
            df_custos_filtrado['Classificacao_Contabil_2'] = df_custos_filtrado['Classificacao_Contabil_2'].replace('MDO PJ Fixo', 'MDO Terceirizada - Artístico')
        else:   
            df_custos_filtrado = df_custos_filtrado[
                (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) |
                (df_custos_filtrado['Classificacao_Contabil_2'] == 'MDO Terceirizada - Artístico') 
            ].copy()

    elif class_cont == 'Custos de Eventos': # Realoca MDO de PJ para Eventos
        df_custos_filtrado = df_custos_filtrado[
            ((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) |
            (df_custos_filtrado['Cargo_DRE'] == 'MDO Terceirizada - Eventos')) &
            (~df_custos_filtrado['Classificacao_Contabil_2'].isin(['  -  Comissões e Gorjeta', 'Conduções/Taxi/Uber', 'Insumos - Alimentos', '(-) Despesas de Patrocínio'])) # Caso - class. cont. 2 erradas em Custos de Eventos
        ].copy()

        # Faz a renomeação por conta da filtragem por Cargo_DRE
        df_custos_filtrado['Classificacao_Contabil_2'] = df_custos_filtrado['Classificacao_Contabil_2'].replace(
            'MDO PJ Fixo',
            'MDO Terceirizada - Eventos'
        )

    elif class_cont == 'Deduções sobre Venda':
        df_custos_filtrado = df_custos_filtrado[
            ((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
            (~df_custos_filtrado['Classificacao_Contabil_2'].isin(['Tarifas Bancárias']))) # Caso - class. cont. 2 errada em Deduções sobre Venda
        ].copy()

    elif class_cont == 'Mão de Obra - Salários':
        df_custos_filtrado = df_custos_filtrado[
            ((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) |
            (df_custos_filtrado['Classificacao_Contabil_2'] == '  -  INSS Segurados') |
            (df_custos_filtrado['Classificacao_Contabil_2'] == 'IRRF - MDO CLT - Salário')) 
        ].copy()

        # Faz a renomeação por conta das class. cont. 2 que foram selecionadas que não são iguais as da class.cont. 1 de salários
        df_custos_filtrado['Classificacao_Contabil_2'] = (df_custos_filtrado['Classificacao_Contabil_2'].replace({
                '  -  INSS Segurados': 'MDO CLT - Salário',
                'IRRF - MDO CLT - Salário': 'MDO CLT - Salário'
            })
        )

    elif class_cont == 'Mão de Obra - PJ':
        if casa == 'Blue Note - São Paulo':
            df_custos_filtrado = df_custos_filtrado[
                (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
                (~df_custos_filtrado['Cargo_DRE'].isin(['MDO Terceirizada - Eventos', '  - Administrativa']) &
                (~df_custos_filtrado['Cargo_DRE'].isna()))
            ].copy()
        if casa == 'Ultra Evil Premium Ltda ':
            df_custos_filtrado = df_custos_filtrado[
                (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
                (~df_custos_filtrado['Cargo_DRE'].isin(['MDO Terceirizada - Eventos', '  - Diretoria', '  - Assistente']) &
                (~df_custos_filtrado['Cargo_DRE'].isna()) |
                (df_custos_filtrado['Cargo_DRE'].isin(['  - Manutenção'])))
            ].copy()
        else:
            df_custos_filtrado = df_custos_filtrado[
                (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
                (~df_custos_filtrado['Cargo_DRE'].isin(['MDO Terceirizada - Eventos', '  - Diretoria', '  - Assistente']) &
                (~df_custos_filtrado['Cargo_DRE'].isna()))
            ].copy()
    
    elif class_cont == 'Mão de Obra - Encargos e Provisões':
        df_custos_filtrado = df_custos_filtrado[
            ((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
            (~df_custos_filtrado['Classificacao_Contabil_2'].isin(['  -  INSS Segurados', 'IRRF - MDO CLT - Salário', 'IRRF']))) # Caso - class. cont. 2 em MDO - Encargos e Provisões não exibidas
        ].copy()

    elif class_cont == 'Mão de Obra - Benefícios':
        df_custos_filtrado = df_custos_filtrado[
            (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
            (df_custos_filtrado['Classificacao_Contabil_2'] != 'Contribuição Sindical Assistencial   ')
        ].copy()

    elif class_cont == 'Utilidades': 
        df_custos_filtrado = df_custos_filtrado[
            (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
            (df_custos_filtrado['Classificacao_Contabil_2'] != 'Custas Cartório / Operação') # Caso - class. cont. 2 em Utilidades não exibidas
        ].copy()

    elif class_cont == 'Serviços de Terceiros': 
        df_custos_filtrado = df_custos_filtrado[
            (df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
            (df_custos_filtrado['Classificacao_Contabil_2'] != 'VALET E MOTOBOY') # Caso - class. cont. 2 em Serviços de Terceiros não exibidas
        ].copy() 

    elif class_cont == 'Despesas Financeiras':
        if casa == 'Arcos':
            df_custos_filtrado = df_custos_filtrado[
            (((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
            (df_custos_filtrado['Classificacao_Contabil_2'] == 'Tarifas Bancárias')) |
            (df_custos_filtrado['Fornecedor'] == 'NELSON WILIANS E ADVOGADOS ASSOCIADOS - MATRIZ'))
        ].copy()
        else:
            df_custos_filtrado = df_custos_filtrado[
                ((df_custos_filtrado['Classificacao_Contabil_1'] == class_cont) &
                (df_custos_filtrado['Classificacao_Contabil_2'] == 'Tarifas Bancárias'))
            ].copy()
        
    elif class_cont == 'Desconto sobre Venda':
        df_descontos_filtrado = df_custos_filtrado[
            (df_custos_filtrado['Descontos - DRE'].isin(['Descontos - Operação', 'Desconto - Alimentação Escritório', 'Descontos - Marketing']))
        ].copy()

        # Renomeia essa coluna para poder aplicar o código abaixo e cria coluna de data
        df_descontos_filtrado = df_descontos_filtrado.rename(columns={'Descontos - DRE': 'Classificacao_Contabil_2'})
        df_descontos_filtrado['Data_Competencia'] = pd.to_datetime(
            dict(
                year=df_descontos_filtrado['Ano'],
                month=df_descontos_filtrado['Mês'],
                day=1
            )
        )
        df_custos_filtrado = df_descontos_filtrado.copy()
    
    # Implementa cálculo de Sistema de Franquias - Fee Gestão FB para casas 100% FB (meses passados para ser possível projetar)
    elif class_cont == 'Sistema de Franquias' and casa not in ['Arcos', 'Blue Note - São Paulo', 'Love Cabaret', 'Ultra Evil Premium Ltda ']:
        df_custos_filtrado = df_valor_fee_gestao.copy()
        df_custos_filtrado['Data_Competencia'] = pd.to_datetime({
            'year': df_custos_filtrado['Ano'],
            'month': df_custos_filtrado['Mês'],
            'day': 1
        }).dt.date        
        df_custos_filtrado = df_custos_filtrado.rename(columns={'Valor Bruto': 'Valor_Pagamento'})
        df_custos_filtrado['Valor_Liquido'] = df_custos_filtrado['Valor_Pagamento']
        
    else:
        df_custos_filtrado = df_custos_filtrado[df_custos_filtrado['Classificacao_Contabil_1'] == class_cont].copy()

    if class_cont == 'Desconto sobre Venda':
        col_valor = ['Permanece no Desconto']
    else:
        col_valor = ['Valor_Pagamento', 'Valor_Liquido']
    
    # Cria colunas de mês e ano e soma o total mensal para cada class. cont. 2
    df_custos_filtrado['Data_Competencia'] = pd.to_datetime(df_custos_filtrado['Data_Competencia'], errors='coerce')
    df_custos_filtrado['Ano'] = df_custos_filtrado['Data_Competencia'].dt.year
    df_custos_filtrado['Mês'] = df_custos_filtrado['Data_Competencia'].dt.month
    df_custos_filtrado_mensal = df_custos_filtrado.groupby(['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2'], as_index=False)[col_valor].sum()
    df_custos_filtrado_mensal = df_custos_filtrado_mensal.rename(columns={col_valor[0]:'Custo Real'})

    # Casos em que o custo mensal não depende apenas da aut_blue_me_sem_pedido: merge com outra tabela
    if class_cont in ['Mão de Obra - Salários', 'Mão de Obra - Benefícios', 'Gorjeta', 'Custos Artístico Geral', 'Marketing', 'Informática e TI', 'Patrocínio', 'Despesas Financeiras']: 
        df_custos_filtrado_mensal = merge_despesas_complexas(df_custos_filtrado_mensal, df_tabela_secundaria, df_tabela_terciaria, df_tabela_quaternaria, casa, class_cont)

    if class_cont == 'Utilidades':
        df_aut_filtrado = df_aut_blue_me_com_pedido[
            (df_aut_blue_me_com_pedido['Casa'] == casa)
        ].copy()
        
        # Cria colunas de mês e ano e soma o total mensal para cada class. cont. 2
        df_aut_filtrado['Data_Emissao'] = pd.to_datetime(df_aut_filtrado['Data_Emissao'], errors='coerce')
        df_aut_filtrado['Ano'] = df_aut_filtrado['Data_Emissao'].dt.year
        df_aut_filtrado['Mês'] = df_aut_filtrado['Data_Emissao'].dt.month
        df_aut_filtrado = df_aut_filtrado.groupby(['Casa', 'Mês', 'Ano'], as_index=False)[['Valor_Liq_Alimentos', 'Valor_Liq_Bebidas', 'Valor_Liq_Descart_Hig_Limp', 'Valor_Gelo_Gas_Carvao_Velas', 'Valor_Utensilios', 'Valor_Liq_Outros']].sum()

        df_custos_base = df_custos_filtrado_mensal[['Casa', 'Classificacao_Contabil_2']].drop_duplicates()
        df_combinado = df_custos_base.merge(df_aut_filtrado[['Ano', 'Mês', 'Valor_Liq_Alimentos', 'Valor_Liq_Bebidas', 'Valor_Liq_Descart_Hig_Limp', 'Valor_Gelo_Gas_Carvao_Velas', 'Valor_Utensilios', 'Valor_Liq_Outros']], how='cross')

        # Agora junta novamente com os custos reais (para pegar valores quando existirem)
        df_merge = pd.merge(
            df_combinado,
            df_custos_filtrado_mensal[['Casa', 'Ano', 'Mês', 'Classificacao_Contabil_2', 'Custo Real']],
            on=['Casa', 'Ano', 'Mês', 'Classificacao_Contabil_2'],
            how='left'
        )

        df_merge = df_merge.fillna(0)
        conditions = [
            df_merge['Classificacao_Contabil_2'] == 'Higiene e Limpeza',
            df_merge['Classificacao_Contabil_2'] == 'Material de Escritorio',
            df_merge['Classificacao_Contabil_2'] == 'Utensilios',
            df_merge['Classificacao_Contabil_2'] == 'Material de Consumo - Gelo/ Gas CO2/ Carvao /Velas'
        ]

        choices = [
            df_merge['Custo Real'] + df_merge['Valor_Liq_Descart_Hig_Limp'],
            df_merge['Custo Real'] + df_merge['Valor_Liq_Outros'],
            df_merge['Custo Real'] + df_merge['Valor_Utensilios'],
            df_merge['Custo Real'] + df_merge['Valor_Gelo_Gas_Carvao_Velas']
        ]

        df_merge['Custo Real'] = np.select(
            conditions,
            choices,
            default=df_merge['Custo Real']  # mantém o valor original se nenhuma condição for atendida
        )
        df_custos_filtrado_mensal = df_merge[['Casa', 'Mês', 'Ano', 'Classificacao_Contabil_2', 'Custo Real']]

    # Resgata faturamentos projetados por mês
    df_faturamento_meses_futuros['Valor Bruto'] = pd.to_numeric(df_faturamento_meses_futuros['Valor Bruto'], errors='coerce')
    df_faturamento_meses_futuros['Valor Projetado'] = pd.to_numeric(df_faturamento_meses_futuros['Valor Projetado'], errors='coerce')
    df_resgata_faturamento_meses_futuros = df_faturamento_meses_futuros.groupby(['Ano', 'Mês'], as_index=False)[['Valor Bruto', 'Valor Projetado']].sum()
    df_resgata_faturamento_meses_futuros = df_resgata_faturamento_meses_futuros.rename(columns={'Valor Bruto':'Faturamento Real', 'Valor Projetado':'Faturamento Projetado'})
    
    # Merge da tabela de custos passados com a de faturamentos - obter combinação de cada class. cont. 2 com todos os meses do ano para projetar
    df_custos = df_custos_filtrado_mensal.copy()
    df_fat = df_resgata_faturamento_meses_futuros.copy()
    
    # Lista com todas class. cont. 2 da class. cont. 1 selecionada (para incluir as que estão no orçamento mas não geraram custo)
    df_orcamentos_da_class_cont_selecionada = df_orcamentos[df_orcamentos['Classificacao_Contabil_1'] == class_cont].copy()
    lista_class_cont_2_orcamento = df_orcamentos_da_class_cont_selecionada['Categoria'].unique().tolist()
    lista_class_cont_2_da_class_cont_1_selecionada = df_custos['Classificacao_Contabil_2'].unique().tolist()
    lista_completa_class_cont_2 = list(
        set(lista_class_cont_2_da_class_cont_1_selecionada) |
        set(lista_class_cont_2_orcamento)
    )
    
    # Pega apenas colunas de identificação de categoria
    df_custos_base = df_custos[['Casa', 'Classificacao_Contabil_2']].drop_duplicates()

    # Faz o produto cartesiano: cada categoria × cada mês/ano
    df_combinado = df_custos_base.merge(df_fat[['Ano', 'Mês', 'Faturamento Real', 'Faturamento Projetado']], how='cross')
    
    # Agora junta novamente com os custos reais (para pegar valores quando existirem)
    df_custos_faturamentos_mensais_passados = pd.merge(
        df_combinado,
        df_custos[['Casa', 'Ano', 'Mês', 'Classificacao_Contabil_2', 'Custo Real']],
        on=['Casa', 'Ano', 'Mês', 'Classificacao_Contabil_2'],
        how='left'
    )

    # Concatena o df de faturamentos com o de orçamentos de cada class. cont. 2
    df_resultante = organiza_despesas_orcamentos(df_custos_faturamentos_mensais_passados, df_orcamentos, casa, lista_completa_class_cont_2, class_cont)
    return df_resultante


# Para exibir corretamente o orçamento mesmo de despesas que não geraram custo
def organiza_despesas_orcamentos(df_custos, df_orcamentos, casa, lista_completa_class_cont_2, class_cont_1):
    df_orcamentos_filtrado = df_orcamentos[df_orcamentos['Casa'] == casa].copy()
    df_resultante = pd.merge(
        df_custos,
        df_orcamentos_filtrado,
        left_on=['Casa', 'Classificacao_Contabil_2', 'Ano', 'Mês'],
        right_on=['Casa', 'Categoria', 'Ano', 'Mês'],
        how='outer'
    )
    df_resultante['Classificacao_Contabil_2'] = df_resultante['Classificacao_Contabil_2'].fillna(df_resultante['Categoria'])
    df_resultante = df_resultante[df_resultante['Classificacao_Contabil_2'].isin(lista_completa_class_cont_2)].copy()

    if class_cont_1 == 'Deduções sobre Venda': df_resultante = df_resultante[~df_resultante['Classificacao_Contabil_2'].isin(['Desconto - Alimentação Escritório', 'Descontos - Marketing', 'Descontos - Operação'])].copy()
    if class_cont_1 == 'Mão de Obra - PJ': df_resultante = df_resultante[df_resultante['Classificacao_Contabil_2'].isin(['MDO PJ Fixo'])].copy()
    if class_cont_1 == 'Utilidades': df_resultante = df_resultante[~df_resultante['Classificacao_Contabil_2'].isin(['Conduções/Taxi/Uber', 'Viagens e Estadias - Operação'])].copy()

    return df_resultante


def projecao_custos_proximos_meses(df_merge_custos_faturamentos_mensais, class_cont_custo, ano_atual, mes_atual):
    # Cria coluna da porcentagem custo/faturamento a ser projetada
    df_merge_custos_faturamentos_mensais['Custo Percentual Projetado'] = None
    df_merge_custos_faturamentos_mensais['Custo Projetado'] = None

    # Cria colunas auxiliares de data
    df_merge_custos_faturamentos_mensais['Mês'] = df_merge_custos_faturamentos_mensais['Mês'].astype(int)
    df_merge_custos_faturamentos_mensais['Mes_Ano'] = (df_merge_custos_faturamentos_mensais['Ano'].astype(int).astype(str) + '-' + df_merge_custos_faturamentos_mensais['Mês'].astype(int).astype(str).str.zfill(2))    
    df_merge_custos_faturamentos_mensais['Data'] = pd.to_datetime(df_merge_custos_faturamentos_mensais['Mes_Ano'], format='%Y-%m')
    
    # Premissa 1: Outras class. cont. de custos (custo1 + custo2 / fat1 + fat2)
    if class_cont_custo not in ['PJ', 'Salários', 'Custo de Ocupação', 'Informática e TI', 'Serviços de Terceiros', 'Locação de Equipamentos', 'Sistema de Franquias']:
        # Loop por classificação contábil 2
        for class_cont in df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'].dropna().unique():
            df_class_cont = df_merge_custos_faturamentos_mensais[df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'] == class_cont].copy()
            
            if df_class_cont is not None and not df_class_cont.empty:
                for i, row in df_class_cont.iterrows():
                    data = row['Data']

                    # pega histórico dos dois meses atrás
                    dois_meses_atras = data - pd.DateOffset(months=2)

                    historico = df_merge_custos_faturamentos_mensais[
                        (df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'] == class_cont) &
                        (df_merge_custos_faturamentos_mensais['Data'] >= dois_meses_atras) &
                        (df_merge_custos_faturamentos_mensais['Data'] < data)
                    ].copy()

                    # Faz Projecao = (Custo1 + Custo2) / (Faturamento_Geral1 + Faturamento_Geral2)
                    # Define colunas auxiliares conforme o mês
                    historico["Custo_Usado"] = np.where(
                        (historico["Mês"] >= mes_atual) & (historico['Ano'] == ano_atual),
                        historico["Custo Projetado"],       # usa o projetado se mês >= atual
                        historico["Custo Real"]             # senão usa o real
                    )

                    historico["Faturamento_Usado"] = np.where(
                        (historico["Mês"] >= mes_atual) & (historico['Ano'] == ano_atual),
                        historico["Faturamento Projetado"],          # usa o projetado se mês >= atual
                        historico["Faturamento Real"]                # senão usa o real
                    )
                    
                    valores_para_soma_custos = historico['Custo_Usado'].fillna(0).astype(float)
                    valores_para_soma_faturamento = historico['Faturamento_Usado'].fillna(historico['Faturamento Projetado']).astype(float)

                    soma_custos = valores_para_soma_custos.sum()
                    soma_faturamentos = valores_para_soma_faturamento.sum()
                    
                    if soma_faturamentos and not pd.isna(soma_faturamentos) and soma_faturamentos != 0:
                        custo_projetado = (soma_custos / soma_faturamentos) * 100
                    else:
                        custo_projetado = 0 

                    # Atribui o valor à coluna correta
                    df_merge_custos_faturamentos_mensais.at[i, 'Custo Percentual Projetado'] = custo_projetado
                    
                    # Define valor de Custo Projetado em Reais - QUESTÃO dos valores arredondados na exibição do 'Custo Percentual Projetado'
                    df_merge_custos_faturamentos_mensais['Custo Projetado'] = (df_merge_custos_faturamentos_mensais['Custo Percentual Projetado'] / 100) * df_merge_custos_faturamentos_mensais['Faturamento Projetado']

    # Premissa 3: 5% do Faturamento Estimado - Sistema de Franquias
    elif class_cont_custo == 'Sistema de Franquias':
        df_merge_custos_faturamentos_mensais['Custo Projetado'] = 0.05 * df_merge_custos_faturamentos_mensais['Faturamento Projetado']

    # Premissa 2: Igual ao mês anterior - PJ, Salários, Custo de Ocupação, Informática e TI, Serviços de Terceiros, Locação de Equipamentos
    else:
        for class_cont in df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'].dropna().unique():
            df_class_cont = df_merge_custos_faturamentos_mensais[df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'] == class_cont]
            for mes_ano in df_class_cont['Mes_Ano'].unique():
                df_mes_ano = df_class_cont[df_class_cont['Mes_Ano'] == mes_ano]
                data = df_mes_ano['Data'].iloc[0]

                # pega o valor do custo do mês anterior
                mes_anterior = data - pd.DateOffset(months=1)
                dado_mes_anterior = df_class_cont[df_class_cont['Data'] == mes_anterior]
                
                if not dado_mes_anterior.empty:
                    mes_anterior = dado_mes_anterior['Data'].iloc[0].month
                    custo_usado = dado_mes_anterior['Custo Real'].iloc[0]

                    if not custo_usado or mes_anterior == mes_atual: # Caso o mês ainda não tenha o valor lançado
                        custo_usado = dado_mes_anterior['Custo Projetado'].iloc[0] # Usa a projeção do mês anterior em vez do real
                    
                    # Atribui o valor à coluna correta
                    df_merge_custos_faturamentos_mensais.loc[
                        (df_merge_custos_faturamentos_mensais['Mes_Ano'] == mes_ano) &
                        (df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'] == class_cont),
                        'Custo Projetado'
                    ] = custo_usado   

                    # Recarrega o DataFrame atualizado para que o próximo loop "veja" as mudanças
                    df_class_cont = df_merge_custos_faturamentos_mensais[
                        df_merge_custos_faturamentos_mensais['Classificacao_Contabil_2'] == class_cont
                    ]

    return df_merge_custos_faturamentos_mensais


############################################ PREPARA DESPESAS - POR CLASS. CONT. ############################################
def filtra_despesas_mes_ano_selecionados(df, mes, ano):
    df_filtrado = df[
        (df['Ano'] == ano) &
        (df['Mês'] == mes)
    ].copy()
    return df_filtrado


def padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, categoria):
    if categoria in ['  -  Adicional Noturno', '  -  Férias', '  -  13º Salário']: # Incidência: Gorjeta e Salários
        df_categoria = pd.merge(df_gorjeta, df_salarios, on=['Mês', 'Ano'], how='left')
        df_categoria['Custo Gorjeta'] = df_categoria['Custo Gorjeta'].fillna(0)
        df_categoria['Custo Salários'] = df_categoria['Custo Salários'].fillna(0)
        df_categoria['Custo Gorjeta'] = pd.to_numeric(df_categoria['Custo Gorjeta'], errors='coerce')
        df_categoria['Custo Salários'] = pd.to_numeric(df_categoria['Custo Salários'], errors='coerce')

        df_categoria[categoria] = df_categoria['Custo Gorjeta'] + df_categoria['Custo Salários'] 
        df_categoria[categoria] = pd.to_numeric(df_categoria[categoria]).astype(float)

        # Cálculo específico de cada categoria (hardcoded)
        if categoria == '  -  Adicional Noturno': df_categoria[categoria] = (df_categoria[categoria] / 220) * (5*1.14)
        elif categoria == '  -  Férias': df_categoria[categoria] = (df_categoria[categoria] * 0.333) / 12

        df_categoria['Classificacao_Contabil_2'] = categoria
        df_mdo = pd.merge(df_mdo, df_categoria[['Mês', 'Ano', 'Classificacao_Contabil_2', categoria]], on=['Mês', 'Ano', 'Classificacao_Contabil_2'], how='left')
        df_mdo['Custo Real'] = pd.to_numeric(df_mdo['Custo Real']).astype(float)
        df_mdo['Taxa'] = pd.to_numeric(df_mdo['Taxa']).astype(float)

        if categoria in ['  -  Adicional Noturno', '  -  Férias']:
            df_mdo.loc[df_mdo['Classificacao_Contabil_2'] == categoria, 'Custo Real'] = df_mdo[categoria]
        elif categoria == '  -  13º Salário': # 13º Salário é dividido por 12 e multiplicado pela taxa
            df_mdo.loc[df_mdo['Classificacao_Contabil_2'] == categoria, 'Custo Real'] = (df_mdo[categoria] / 12) * (df_mdo['Taxa'] / 100)
        
        df_mdo = df_mdo.drop(columns=[categoria])

    elif categoria in ['  -  FGTS', '  -  INSS Patronal', '  -  INSS Segurados']: # Incidêcia: Gorjeta, Salários, Adicional Noturno
        df_adicional_noturno = df_mdo[df_mdo['Classificacao_Contabil_2'] == '  -  Adicional Noturno'].copy()
        df_adicional_noturno = df_adicional_noturno[['Classificacao_Contabil_2', 'Mês', 'Ano', 'Custo Real']]
        df_adicional_noturno = df_adicional_noturno.rename(columns={'Custo Real': 'Custo Adicional Noturno'})
        df_categoria = df_gorjeta.merge(df_salarios, on=['Mês', 'Ano'], how='left').merge(df_adicional_noturno, on=['Mês', 'Ano'], how='left')
        
        df_categoria['Custo Gorjeta'] = pd.to_numeric(df_categoria['Custo Gorjeta']).astype(float).fillna(0)
        df_categoria['Custo Salários'] = pd.to_numeric(df_categoria['Custo Salários']).astype(float).fillna(0)
        df_categoria['Custo Adicional Noturno'] = pd.to_numeric(df_categoria['Custo Adicional Noturno']).astype(float).fillna(0)
        df_categoria[categoria] = df_categoria['Custo Gorjeta'] + df_categoria['Custo Salários'] + df_categoria['Custo Adicional Noturno']
        df_categoria['Classificacao_Contabil_2'] = categoria
        df_mdo = pd.merge(df_mdo, df_categoria[['Mês', 'Ano', 'Classificacao_Contabil_2', categoria]], on=['Mês', 'Ano', 'Classificacao_Contabil_2'], how='left')
        
        df_mdo['Custo Real'] = pd.to_numeric(df_mdo['Custo Real']).astype(float)
        df_mdo[categoria] = pd.to_numeric(df_mdo[categoria]).astype(float)
        df_mdo['Taxa'] = pd.to_numeric(df_mdo['Taxa']).astype(float)
        df_mdo.loc[df_mdo['Classificacao_Contabil_2'] == categoria, 'Custo Real'] = df_mdo[categoria] * (df_mdo['Taxa'] / 100)

    elif categoria in ['  -  FGTS sobre 13º Salários', '  -  INSS sobre 13º Salários', '  -  INSS sobre Férias']: # Incidência: 13º Salário ou Férias
        if categoria in ['  -  FGTS sobre 13º Salários', '  -  INSS sobre 13º Salários']:
            df_categoria = df_mdo[df_mdo['Classificacao_Contabil_2'] == '  -  13º Salário'].copy()
        elif categoria == '  -  INSS sobre Férias':
            df_categoria = df_mdo[df_mdo['Classificacao_Contabil_2'] == '  -  Férias'].copy()
        
        df_categoria = df_categoria[['Classificacao_Contabil_2', 'Mês', 'Ano', 'Custo Real']]
        df_categoria = df_categoria.rename(columns={'Custo Real': categoria})
        df_categoria['Classificacao_Contabil_2'] = categoria
        
        df_mdo = pd.merge(df_mdo, df_categoria, on=['Mês', 'Ano', 'Classificacao_Contabil_2'], how='left')
        df_mdo.loc[df_mdo['Classificacao_Contabil_2'] == categoria, 'Custo Real'] = df_mdo[categoria] * (df_mdo['Taxa'] / 100)
        
    return df_mdo
        

def calculo_mdo_encargos_provisoes(df_despesas_mdo, df_gorjeta, df_salarios, df_parametros_mdo, casa):    
    # Aplica taxas de acordo com as faixas de data de vigência
    df_despesas_mdo['Data'] = pd.to_datetime({
        'year': df_despesas_mdo['Ano'],
        'month': df_despesas_mdo['Mês'],
        'day': 1
    })
    df_mdo = df_despesas_mdo.merge(df_parametros_mdo[['Data Inicio', 'Data Fim', 'Taxa', 'Classificacao_Contabil_2']], how='cross')
    df_mdo = df_mdo[
        ((df_mdo['Data Fim'].notna()) &
        (df_mdo['Data'] >= df_mdo['Data Inicio']) &
        (df_mdo['Data'] <= df_mdo['Data Fim']) &
        (df_mdo['Classificacao_Contabil_2_x'] == df_mdo['Classificacao_Contabil_2_y']))
        |
        ((df_mdo['Data Fim'].isna()) & # Caso a data de fim de vigência seja nula (vigência atual)
        (df_mdo['Data'] >= df_mdo['Data Inicio']) &
        (df_mdo['Classificacao_Contabil_2_x'] == df_mdo['Classificacao_Contabil_2_y']))]

    df_mdo = df_mdo.drop(columns=['Classificacao_Contabil_2_y', 'Data', 'Data Inicio', 'Data Fim'])
    df_mdo = df_mdo.rename(columns={'Classificacao_Contabil_2_x': 'Classificacao_Contabil_2'})
    df_mdo['Custo Real'] = df_mdo['Custo Real'].fillna(0)

    # Prepara df de gorjeta e salários
    df_gorjeta = df_gorjeta[['Classificacao_Contabil_1', 'Mês', 'Ano', 'Custo Real']]
    df_gorjeta = df_gorjeta.rename(columns={'Custo Real': 'Custo Gorjeta'})
    df_salarios = df_salarios[['Classificacao_Contabil_1', 'Mês', 'Ano', 'Custo Real']]
    df_salarios = df_salarios.rename(columns={'Custo Real': 'Custo Salários'})

    # Aplica a lógica de cada casa
    # Se aplicam a todas as casas
    if casa in ['Arcos', 'Bar Brahma - Centro', 'Bar Brahma - Granja', 'Bar Brahma - Paulista', 'Bar Léo - Centro', 'Blue Note - São Paulo', 'Jacaré', 'Love Cabaret', 'Orfeu', 'Riviera Bar', 'Priceless', 'Ultra Evil Premium Ltda ']: 
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  13º Salário')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  FGTS sobre 13º Salários')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  INSS sobre 13º Salários')
    if casa in ['Bar Brahma - Centro']: # Apenas BBC
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  Adicional Noturno')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  Férias')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  FGTS')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  INSS Patronal')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  INSS Segurados')
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  INSS sobre Férias')
    if casa in ['Bar Brahma - Granja']: # Apenas BBG
        df_mdo = padroes_calculo_mdo_encargos_provisoes(df_mdo, df_gorjeta, df_salarios, '  -  Adicional Noturno')

    return df_mdo


def loop_prepara_dados_despesas(lista_categorias_despesas, df_descontos, df_consumo_interno_cmv, df_consumo_cartao_black, df_aut_blueme_sem_pedido, df_aut_blue_me_com_pedido, df_faturamento_meses_futuros, df_aut_folha, df_orcamentos, df_demais_receitas_extr, df_ajustes_manuais, df_parametros_mdo, df_valor_fee_gestao, df_resultados, casa, mes_selecionado, ano_selecionado):
    for categoria_despesa in lista_categorias_despesas:
        # Define df de despesas utilizado pela categoria
        if categoria_despesa == 'Desconto sobre Venda':
            df_despesas = df_descontos.copy()
        else:
            df_despesas = df_aut_blueme_sem_pedido.copy()
        
        # Define df de despesas complementares utilizado pela categoria
        if categoria_despesa in ['Custos Artístico Geral', 'Marketing', 'Informática e TI', 'Mão de Obra - Benefícios']:
            df_tabela_secundaria = df_descontos.copy()
            df_tabela_terciaria = df_consumo_interno_cmv.copy()
            df_tabela_quaternaria = df_consumo_cartao_black.copy()
        elif categoria_despesa == 'Gorjeta' or categoria_despesa == 'Mão de Obra - Salários':
            df_tabela_secundaria = df_aut_folha.copy()
            df_tabela_terciaria = None
            df_tabela_quaternaria = None
        elif categoria_despesa == 'Patrocínio':
            df_tabela_secundaria = df_demais_receitas_extr[df_demais_receitas_extr['Classificacao'] == 'Patrocínio'].copy()
            if casa == 'Blue Note - São Paulo': 
                df_tabela_terciaria = df_descontos[df_descontos['Descontos - DRE'] == 'Descontos - Operação'].copy()
            else:
                df_tabela_terciaria = None
            df_tabela_quaternaria = None
        elif categoria_despesa == 'Despesas Financeiras': 
            if casa == 'Blue Note - São Paulo': 
                df_tabela_secundaria = df_demais_receitas_extr[df_demais_receitas_extr['Classificacao'].isin(['Cashback', 'Bilheteria-Rebate'])].copy()
            else:
                df_tabela_secundaria = df_aut_blueme_sem_pedido.copy()
            df_tabela_terciaria = None
            df_tabela_quaternaria = None
        else: 
            df_tabela_secundaria = None
            df_tabela_terciaria = None
            df_tabela_quaternaria = None

        # Utilidades utiliza também dados de blueme com pedido
        if categoria_despesa == 'Utilidades':
            df_aut_blueme_com_pedido = df_aut_blue_me_com_pedido.copy()
        else:
            df_aut_blueme_com_pedido = None
        
        df_despesas_mensais_passadas = prepara_dados_custos_mensais(
            df_despesas, 
            df_faturamento_meses_futuros, 
            casa, 
            categoria_despesa, 
            df_orcamentos, 
            df_valor_fee_gestao, # Implementa cálculo de Sistema de Franquias - Fee Gestão FB para casas 100% FB
            df_aut_blue_me_com_pedido=df_aut_blueme_com_pedido, 
            df_tabela_secundaria=df_tabela_secundaria, 
            df_tabela_terciaria=df_tabela_terciaria, 
            df_tabela_quaternaria=df_tabela_quaternaria
        )

        # Cálculo - Encargos e Provisões
        if categoria_despesa == 'Gorjeta':
            df_gorjeta = df_despesas_mensais_passadas.copy()
        if categoria_despesa == 'Mão de Obra - Salários':
            df_salarios = df_despesas_mensais_passadas.copy()
        if categoria_despesa == 'Mão de Obra - Encargos e Provisões': 
            lista_categorias_calculadas = df_parametros_mdo['Classificacao_Contabil_2'].unique().tolist()
            df_despesas_mensais_blueme = df_despesas_mensais_passadas[~df_despesas_mensais_passadas['Classificacao_Contabil_2'].isin(lista_categorias_calculadas)]
            df_despesas_mensais_calculadas = calculo_mdo_encargos_provisoes(df_despesas_mensais_passadas, df_gorjeta, df_salarios, df_parametros_mdo, casa)
            df_despesas_mensais_passadas = pd.concat([df_despesas_mensais_blueme, df_despesas_mensais_calculadas])

        # Merge com ajustes manuais para despesas que tem lançamento de ajuste
        df_ajustes_categoria = df_ajustes_manuais[
            (df_ajustes_manuais['Casa'] == casa) &
            (df_ajustes_manuais['Ano'] == ano_selecionado) &
            (df_ajustes_manuais['Classificacao_Contabil_1'] == categoria_despesa)
        ].copy()
        
        df_ajustes_categoria = df_ajustes_categoria.groupby(['ID_Casa', 'Casa', 'Mês', 'Ano', 'Classificacao_Contabil_1', 'Classificacao_Contabil_2'], as_index=False)['Valor Ajuste'].sum()

        if not df_ajustes_categoria.empty:
            df_despesas_mensais_passadas['Classificacao_Contabil_1'] = df_despesas_mensais_passadas['Classificacao_Contabil_1'].fillna(categoria_despesa)
            df_despesas_com_ajustes = pd.merge(
                df_despesas_mensais_passadas,
                df_ajustes_categoria,
                on=['Classificacao_Contabil_1', 'Classificacao_Contabil_2', 'Casa', 'Ano', 'Mês'],
                how='left'
            ).fillna(0)
            
            df_despesas_com_ajustes['Valor Ajuste'] = pd.to_numeric(df_despesas_com_ajustes['Valor Ajuste'], errors='coerce')
            df_despesas_com_ajustes['Custo Real'] = pd.to_numeric(df_despesas_com_ajustes['Custo Real'], errors='coerce')

            condicao_subtrair = df_despesas_com_ajustes['Classificacao_Contabil_2'] != '(+) Receitas de Patrocínio' # Caso específico de "despesa" positiva
            df_despesas_com_ajustes.loc[condicao_subtrair, 'Custo Real'] -= (df_despesas_com_ajustes.loc[condicao_subtrair, 'Valor Ajuste']) 

            condicao_somar = ~condicao_subtrair
            df_despesas_com_ajustes.loc[condicao_somar, 'Custo Real'] += (df_despesas_com_ajustes.loc[condicao_somar, 'Valor Ajuste']) 
        
        else: # Se não tem ajuste, mantém as despesas originais
            df_despesas_com_ajustes = df_despesas_mensais_passadas.copy()

        df_projecao_despesa = projecao_custos_proximos_meses(df_despesas_com_ajustes, categoria_despesa, datas['ano_atual'], datas['mes_atual'])
        if categoria_despesa == 'Gorjeta': # Guarda para utilizar no cálculo do Imposto Simples
            df_gorjeta = df_projecao_despesa.copy()
        if categoria_despesa == 'Mão de Obra - Salários':
            df_salarios = df_projecao_despesa.copy()

        df_projecao_despesa = filtra_despesas_mes_ano_selecionados(df_projecao_despesa, mes_selecionado, ano_selecionado)
        df_projecao_despesa = calcula_linha_total(df_projecao_despesa, 'Classificacao_Contabil_2', categoria_despesa, 'Custo Projetado', 'Custo Real')
        df_resultados.append(df_projecao_despesa)

    return df_resultados, df_gorjeta, df_salarios


############################################ CRIAÇÃO LAYOUT E ESTILOS - DRE ############################################
def aplica_layout_dre(df_faturamento_meses_passados_futuros, df_layout_impostos_venda, df_layout_impostos_renda, df_cmv_projetado, df_projecao_despesas, mes_selecionado, ano_selecionado):
    # Formata dados de faturamento
    df_layout_faturamento = df_faturamento_meses_passados_futuros[
        (df_faturamento_meses_passados_futuros['Ano'] == ano_selecionado) &
        (df_faturamento_meses_passados_futuros['Mês'] == mes_selecionado) &
        (df_faturamento_meses_passados_futuros['Categoria'].isin(['Alimentos', 'Bebidas', 'Couvert', 'Serviço', 'Gifts', 'Eventos A&B', 'Eventos Couvert', 'Eventos Locações', 'Eventos Rebate Fornecedores', 'Membership', 'Delivery', 'Outras Receitas']))
    ].copy()

    df_layout_faturamento = df_layout_faturamento.rename(columns={'Valor Bruto': 'Valor Real', 'Atingimento Real': 'Percentual Real (do Orçamento)', 'Projeção Atingimento': 'Percentual Projetado'})
    df_layout_faturamento = calcula_linha_total(df_layout_faturamento, 'Categoria', 'Faturamento', 'Valor Projetado', 'Valor Real')
    df_layout_faturamento = df_layout_faturamento[['Categoria', 'Orçamento', 'Percentual Projetado', 'Valor Projetado', 'Valor Real', 'Percentual Real (do Orçamento)']]

    # Formata dados de CMV
    df_layout_cmv = df_cmv_projetado[(df_cmv_projetado['Ano'] == ano_selecionado) & (df_cmv_projetado['Mês'] == mes_selecionado)]
    df_layout_cmv = df_layout_cmv.drop(columns=['Valor Projetado'])
    df_layout_cmv = df_layout_cmv.rename(columns={'CMV Percentual Projetado': 'Percentual Projetado', 'CMV Projetado': 'Valor Projetado', 'CMV Real': 'Valor Real', 'CMV Real Percentual': 'Percentual Real (do Orçamento)', 'CMV Orçado': 'Orçamento'})
    df_layout_cmv['Categoria'] = 'CMV'
    df_layout_cmv = calcula_linha_total(df_layout_cmv, 'Categoria', 'Custo Mercadoria Vendida', 'Valor Projetado', 'Valor Real')
    df_layout_cmv = df_layout_cmv[['Categoria', 'Orçamento', 'Percentual Projetado', 'Valor Projetado', 'Valor Real', 'Percentual Real (do Orçamento)']]

    # Formata dados de despesas
    df_layout_despesas = df_projecao_despesas.copy()
    df_layout_despesas = df_layout_despesas.drop(columns=['Categoria'])
    df_layout_despesas = df_layout_despesas.rename(columns={'Classificacao_Contabil_2': 'Categoria', 'Custo Percentual Projetado': 'Percentual Projetado', 'Custo Projetado': 'Valor Projetado', 'Custo Real': 'Valor Real'})
    df_layout_despesas = df_layout_despesas[['Categoria', 'Orçamento', 'Percentual Projetado', 'Valor Projetado', 'Valor Real']]

    # --------- Inserções de itens calculados previamente --------- #
    # Impostos sobre Venda - depois de 'Descontos sobre Venda'
    indice = df_layout_despesas[df_layout_despesas['Categoria'] == 'Descontos - Operação'].index.max()
    df_parte1 = df_layout_despesas.loc[:indice]
    df_parte2 = df_layout_despesas.loc[indice+1:]

    df_layout_despesas_final = pd.concat([
        df_parte1,
        df_layout_impostos_venda,
        df_parte2
    ]).reset_index(drop=True)

    # CMV - depois de 'Impostos sobre Venda'
    indice = df_layout_despesas_final[df_layout_despesas_final['Categoria'] == 'ISS'].index.max()
    df_parte1 = df_layout_despesas_final.loc[:indice]
    df_parte2 = df_layout_despesas_final.loc[indice+1:]

    df_layout_despesas_final = pd.concat([
        df_parte1,
        df_layout_cmv,
        df_parte2
    ]).reset_index(drop=True)

    # Impostos de Renda - depois de (-) Despesas de Patrocínio
    indice = df_layout_despesas_final[df_layout_despesas_final['Categoria'] == '(-) Despesas de Patrocínio'].index.max()
    df_parte1 = df_layout_despesas_final.loc[:indice]
    df_parte2 = df_layout_despesas_final.loc[indice+1:]

    df_layout_despesas_final = pd.concat([
        df_parte1,
        df_layout_impostos_renda,
        df_parte2
    ]).reset_index(drop=True)

    # Calcula coluna de Percentual Real 
    df_layout_despesas_final['Orçamento'] = pd.to_numeric(df_layout_despesas_final['Orçamento'], errors='coerce')
    df_layout_despesas_final['Percentual Real (do Orçamento)'] = (df_layout_despesas_final['Valor Real'] / df_layout_despesas_final['Orçamento'].replace(0, np.nan)) * 100

    # Despesas são consideradas negativas
    df_layout_despesas_final.loc[df_layout_despesas_final['Categoria'] != '(+) Receitas de Patrocínio', 'Orçamento'] *= -1    
    df_layout_despesas_final.loc[df_layout_despesas_final['Categoria'] != '(+) Receitas de Patrocínio', 'Valor Projetado'] *= -1    
    df_layout_despesas_final.loc[df_layout_despesas_final['Categoria'] != '(+) Receitas de Patrocínio', 'Valor Real'] *= -1   

    # Concatena os dados
    df_layout_dre = pd.concat([df_layout_faturamento, df_layout_despesas_final])
    df_layout_dre['Orçamento'] = df_layout_dre['Orçamento'].fillna(0)
    df_layout_dre['Percentual Projetado'] = df_layout_dre['Percentual Projetado'].fillna(0)
    df_layout_dre['Valor Projetado'] = pd.to_numeric(df_layout_dre['Valor Projetado'], errors='coerce')
    df_layout_dre['Valor Real'] = pd.to_numeric(df_layout_dre['Valor Real'], errors='coerce')
    df_layout_dre['Percentual Projetado'] = pd.to_numeric(df_layout_dre['Percentual Projetado'], errors='coerce')
    df_layout_dre['Percentual Real (do Orçamento)'] = pd.to_numeric(df_layout_dre['Percentual Real (do Orçamento)'], errors='coerce')

    return df_layout_dre


def calcula_linha_total(df, col_categoria, categoria, col_valor_projetado, col_valor_real):
    df['Orçamento'] = pd.to_numeric(df['Orçamento'], errors='coerce').fillna(0)
    df[col_valor_projetado] = pd.to_numeric(df[col_valor_projetado], errors='coerce').fillna(0)
    df[col_valor_real] = pd.to_numeric(df[col_valor_real], errors='coerce').fillna(0)

    nova_linha = df[['Orçamento', col_valor_projetado, col_valor_real]].sum().to_frame().T
    nova_linha[col_categoria] = categoria

    df = pd.concat([nova_linha, df], ignore_index=True)
    return df


# Função auxiliar para definir linhas calculadas
def soma_categorias(df, categorias, colunas_valores):
    return df[df['Categoria'].isin(categorias)][colunas_valores].sum()


# Calcula porcentagens e outros valores - Orçamento e Real DRE
def define_linhas_calculadas(df_dre, colunas_valores, lista_categorias_despesas, mapa_insercao):
    df_final = df_dre.copy()

    # Define valores mais usados
    cmv = df_final[df_final['Categoria'] == 'Custo Mercadoria Vendida'][colunas_valores].sum()
    custos_artistico = df_final[df_final['Categoria'] == 'Custos Artístico Geral'][colunas_valores].sum()
    faturamento_artistico = df_final[df_final['Categoria'] == 'Couvert'][colunas_valores].sum() # Artístico (couvert/shows)
    faturamento_bruto = df_final[df_final['Categoria'] == 'Faturamento'][colunas_valores].sum()
    custos_eventos = df_final[df_final['Categoria'] == 'Custos de Eventos'][colunas_valores].sum()

    # RECEITA LIQUIDA
    receita_liquida = soma_categorias(df_final, ['Faturamento', 'Desconto sobre Venda', 'Impostos sobre Venda'], colunas_valores)
    df_final = insere_nova_linha(df_final, colunas_valores, receita_liquida, mapa_insercao['RECEITA LÍQUIDA'], 'Categoria', 'RECEITA LÍQUIDA')

    # % sobre Receita Bruta - CMV
    receita_bruta = soma_categorias(df_final, ['Alimentos', 'Bebidas', 'Eventos A&B', 'Delivery'], colunas_valores)
    porc_receita_bruta_cmv = (cmv / receita_bruta)
    df_final = insere_nova_linha(df_final, colunas_valores, porc_receita_bruta_cmv, 'CMV', 'Categoria', '% sobre Receita Bruta')
    
    # % sobre Receita Líquida - CMV
    porc_receita_liquida_cmv = (cmv / receita_liquida).round(2)
    df_final = insere_nova_linha(df_final, colunas_valores, porc_receita_liquida_cmv, '% sobre Receita Bruta', 'Categoria', '% sobre Receita Líquida')

    # % sobre Receita Artístico
    porc_receita_artistico = custos_artistico.div(faturamento_artistico.replace(0, np.nan)).fillna(0).round(2)
    df_final = insere_nova_linha(df_final, colunas_valores, porc_receita_artistico, mapa_insercao['Custos Artístico Geral'], 'Categoria', '% sobre Receita Artístico')

    # % sobre Receita de Eventos
    faturamento_eventos = soma_categorias(df_final, ['Eventos A&B', 'Eventos Locações', 'Eventos Couvert'], colunas_valores)
    porc_receita_eventos = (custos_eventos / faturamento_eventos.replace(0, np.nan)).round(2)
    df_final = insere_nova_linha(df_final, colunas_valores, porc_receita_eventos, mapa_insercao['Custos de Eventos'], 'Categoria', '% sobre Receita de Eventos')

    # MARGEM BRUTA DE CONTRIBUIÇÃO
    margem_bruta_contribuicao = soma_categorias(
        df_final, 
        ['RECEITA LÍQUIDA', 'Deduções sobre Venda', 'Gorjeta', 'Custos de Eventos', 'Custos Artístico Geral', 'Custo Mercadoria Vendida'], 
        colunas_valores
    )
    df_final = insere_nova_linha(df_final, colunas_valores, margem_bruta_contribuicao, mapa_insercao['Deduções sobre Venda'], 'Categoria', 'MARGEM BRUTA DE CONTRIBUIÇÃO')
    lista_categorias_despesas.append('MARGEM BRUTA DE CONTRIBUIÇÃO')

    # PESSOAL
    pessoal = soma_categorias(
        df_final,
        ['Mão de Obra - PJ', 'Mão de Obra - Salários', 'Mão de Obra - Extra', 'Mão de Obra - Encargos e Provisões', 'Mão de Obra - Benefícios'],
        colunas_valores
    )
    df_final = insere_nova_linha(df_final, colunas_valores, pessoal, 'MARGEM BRUTA DE CONTRIBUIÇÃO', 'Categoria', 'PESSOAL')
    lista_categorias_despesas.append('PESSOAL')

    # TOTAL - DESPESAS OPERATIVAS
    total_despesas_operativas = soma_categorias(
        df_final,
        ['PESSOAL', 'Custo de Ocupação', 'Utilidades', 'Informática e TI', 'Manutenção', 'Marketing', 'Serviços de Terceiros', 'Locação de Equipamentos', 'Sistema de Franquias'],
        colunas_valores
    )
    df_final = insere_nova_linha(df_final, colunas_valores, total_despesas_operativas, mapa_insercao['Sistema de Franquias'], 'Categoria', 'TOTAL - DESPESAS OPERATIVAS')
    lista_categorias_despesas.append('TOTAL - DESPESAS OPERATIVAS')
    
    # EBTIDA e EBIT
    total_despesas_operativas = df_final[df_final['Categoria'] == 'TOTAL - DESPESAS OPERATIVAS'][colunas_valores].sum() 
    margem_bruta_contribuicao = df_final[df_final['Categoria'] == 'MARGEM BRUTA DE CONTRIBUIÇÃO'][colunas_valores].sum() 
    ebitda = margem_bruta_contribuicao + total_despesas_operativas
    df_final = insere_nova_linha(df_final, colunas_valores, ebitda, 'TOTAL - DESPESAS OPERATIVAS', 'Categoria', 'EBITDA')
    lista_categorias_despesas.append('EBITDA')
    
    ebit = ebitda
    df_final = insere_nova_linha(df_final, colunas_valores, ebit, 'EBITDA', 'Categoria', 'EBIT')

    # Resultado Antes do IR
    resultado_antes_ir = soma_categorias(
        df_final,
        ['EBIT', '(+/-) Receitas/Despesas Financeiras', '(-) Despesas de Patrocínio', '(+) Receitas de Patrocínio'],
        colunas_valores
    )
    df_final = insere_nova_linha(df_final, colunas_valores, resultado_antes_ir, '(-) Despesas de Patrocínio', 'Categoria', 'Resultado Antes do IR')

    # Resultado Líquido
    resultado_liquido = soma_categorias(df_final, ['Resultado Antes do IR', '(-) Impostos'], colunas_valores)
    df_final = insere_nova_linha(df_final, colunas_valores, resultado_liquido, 'SIMPLES', 'Categoria', 'Resultado Líquido')
    lista_categorias_despesas.append('Resultado Líquido')

    # Total - Variações s/ Resultado Líquido
    total_variacoes = soma_categorias(df_final, ['(-) CAPEX (Investimentos)', '(+/-) Outras variações no fluxo de caixa'], colunas_valores)
    df_final = insere_nova_linha(df_final, colunas_valores, total_variacoes, mapa_insercao['Dividendos e Remunerações Variáveis'], 'Categoria', 'Total - Variações s/ Resultado Líquido')

    # FCF
    fcf = soma_categorias(df_final, ['Resultado Líquido', 'Total - Variações s/ Resultado Líquido'], colunas_valores)
    df_final = insere_nova_linha(df_final, colunas_valores, fcf, 'Total - Variações s/ Resultado Líquido', 'Categoria', 'FCF')

    # Calcula % sobre Receita Bruta de cada categoria
    for categoria in lista_categorias_despesas:
        if categoria not in [ # Casos específicos (não pedem o cálculo)
            'Custo Mercadoria Vendida', 'Impostos sobre Venda', 'Mão de Obra - PJ', 'Mão de Obra - Salários', 'Mão de Obra - Extra', 
            'Mão de Obra - Encargos e Provisões', 'Mão de Obra - Benefícios', 'Patrocínio', 'Despesas Financeiras', 
            'Investimento - CAPEX', 'Dividendos e Remunerações Variáveis', 'Endividamento'
            ]:
            custos_categoria = df_final[df_final['Categoria'] == categoria][colunas_valores].sum()
            porc_faturamento_bruto_categoria = (custos_categoria / faturamento_bruto).round(2)
            if categoria in ['MARGEM BRUTA DE CONTRIBUIÇÃO', 'TOTAL - DESPESAS OPERATIVAS', 'EBITDA', 'Resultado Líquido']:
                apos_linha = categoria
            else:
                apos_linha = mapa_insercao.get(categoria, categoria)
            df_final = insere_nova_linha(df_final, colunas_valores, porc_faturamento_bruto_categoria, apos_linha, 'Categoria', '% sobre Receita Bruta')

    df_final = df_final.fillna(0)
    return df_final


def formatar_colunas_moeda_br(valor):
    if pd.isna(valor) or valor == 0:
        return "-"
    else:
        return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def formatar_colunas_porcentagem(valor):
    if pd.isna(valor) or valor == 0:
        return "-"
    else:
        if valor < 0:
            valor *= (-1)
        return f"{valor:,.2f}%".replace(".", ",")
    

def formatar_linhas_porcentagem(valor):
    if pd.isna(valor) or valor == 0:
        return "-"
    else:
        return f"{valor*100:,.2f}%".replace(".", ",")
    

# Exportação do Excel
def export_to_excel(wb, df, sheet_name):
    if sheet_name in wb.sheetnames:
        wb.remove(wb[sheet_name])
    ws = wb.create_sheet(title=sheet_name)
    
    # Escrever os cabeçalhos
    for col_idx, column_title in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=column_title)
    
    # Escrever os dados
    for row_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


